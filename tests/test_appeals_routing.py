"""Offline Appeals routing regressions; no DB, SDK, GPU or model calls."""
from __future__ import annotations

import ast
import asyncio
import importlib
import importlib.util
import json
import sys
import tempfile
from types import ModuleType, SimpleNamespace
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, Mock, patch

import pandas as pd


SKILL = Path(__file__).resolve().parents[1] / "workspace/skills/appeals-analyzer"
PACKAGE = "appeals_routing_contract"
if PACKAGE not in sys.modules:
    spec = importlib.util.spec_from_file_location(PACKAGE, SKILL / "__init__.py", submodule_search_locations=[str(SKILL)])
    package = importlib.util.module_from_spec(spec)
    sys.modules[PACKAGE] = package
    spec.loader.exec_module(package)

reports = importlib.import_module(f"{PACKAGE}.scripts.appeals_reports")
router = importlib.import_module(f"{PACKAGE}.utils.intent_router")
parser = importlib.import_module(f"{PACKAGE}.utils.greenplum_engine")
sessions = importlib.import_module(f"{PACKAGE}.utils.session_extract_manager")

CANONICAL = "Анализ обращений\n\nПродукт:\n- Банковская карта\n\nПодпродукт:\n- Кредитная карта\n\nЗапрос:\nНайди жалобы на кредитные карты"
INVALID = CANONICAL.replace("Подпродукт:", "Субпродукт:")
JSON_PROMPT = json.dumps({"request_type": "appeals_analysis", "filters": {
    "prd": ["Банковская карта"], "s_prd": ["Кредитная карта"], "chnl": [],
    "date_from": None, "date_to": None}, "prompt": "Найди жалобы на кредитные карты"}, ensure_ascii=False)
CSV_PROMPT = '"Кредиты", "", "", "Найди жалобы на кредитные карты"'
INVALID_JSON = '{"request_type":"appeals_analysis","filters":{},"prompt":"Найди жалобы"}'
NO_SELECTION = "Нет сохранённой релевантной выборки"


class RoutingTests(unittest.TestCase):
    def setUp(self):
        self.keys = ("routing-a", "routing-b")
        for key in self.keys:
            sessions.clear_session_extract(key)
        self.frame = pd.DataFrame({"id": ["123456"], "score": [0.9]})
        self.real_search = reports._search_population
        self.real_extract_dates = reports.extract_search_params
        self.real_export = reports.export_complaints_excel
        patches = {
            "fetch_candidate_ids_by_product": Mock(return_value=["123456"]),
            "_search_population": AsyncMock(return_value=self.frame),
            "export_complaints_excel": Mock(return_value={"name": "mock.xlsx"}),
            "generate_complaint_hypothesis_narrative": AsyncMock(return_value="Готовый отчёт"),
            "extract_search_params": Mock(return_value={"date_range": None}),
            "def_ask_gigachat": Mock(side_effect=AssertionError("No LLM in routing")),
        }
        self.mocks = patches
        for name, replacement in patches.items():
            patcher = patch.object(reports, name, replacement)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.addCleanup(lambda: [sessions.clear_session_extract(key) for key in self.keys])

    def run_report(self, prompt, key="routing-a"):
        return asyncio.run(reports.run_appeals_report(key, prompt))

    def selection(self, key="routing-a"):
        return sessions.set_session_extract(key, self.frame, extra={"final_ids": ["123456"], "hypothesis": "report"})

    def assert_format_error(self, prompt=INVALID):
        result = self.run_report(prompt)
        self.assertNotIn(NO_SELECTION, result)
        self.assertTrue("canonical" in result or "filters" in result or "CSV" in result, result)
        self.mocks["_search_population"].assert_not_awaited()
        return result

    def test_invalid_canonical_empty_session(self):
        self.assertIn("Субпродукт:", self.assert_format_error())

    def test_repeat_invalid_canonical_same_session(self):
        self.assert_format_error()
        self.assert_format_error()

    def test_valid_canonical_after_two_errors(self):
        self.assert_format_error()
        self.assert_format_error()
        self.assertIn("Готовый отчёт", self.run_report(CANONICAL))
        self.assertEqual(sessions.get_session_extract(self.keys[0])["final_ids"], ["123456"])

    def test_invalid_canonical_with_existing_selection(self):
        self.selection()
        with patch.object(reports, "_run_follow_up", new=AsyncMock(return_value="wrong followup")) as followup:
            self.assert_format_error()
            followup.assert_not_awaited()

    def test_valid_new_canonical_with_existing_selection(self):
        self.selection()
        with patch.object(reports, "_run_follow_up", new=AsyncMock()) as followup:
            self.assertIn("Готовый отчёт", self.run_report(CANONICAL))
            followup.assert_not_awaited()

    def test_valid_json_after_invalid_canonical(self):
        self.assert_format_error()
        self.assertIn("Готовый отчёт", self.run_report(JSON_PROMPT))

    def test_valid_csv_after_invalid_json(self):
        self.assert_format_error(INVALID_JSON)
        self.assertIn("Готовый отчёт", self.run_report(CSV_PROMPT))

    def test_followup_existing_selection(self):
        saved = self.selection()
        for prompt in ("Покажи среди этих обращений блокировки карты", "Какие основные причины жалоб среди них?"):
            with patch.object(reports, "_run_follow_up", new=AsyncMock(return_value="followup")) as followup:
                self.assertEqual(self.run_report(prompt), "followup")
                followup.assert_awaited_once_with(self.keys[0], prompt, saved)
        self.mocks["_search_population"].assert_not_awaited()

    def test_no_ids_never_runs_followup_but_allows_natural_search(self):
        for saved in (None, {}, {"final_ids": []}, {"final_ids": None}, {"hypothesis": "technical"}):
            with self.subTest(saved=saved), patch.object(reports, "get_session_extract", return_value=saved), patch.object(reports, "_run_follow_up", new=AsyncMock()) as followup:
                result = self.run_report("Покажи обращения, связанные с блокировкой карты")
                self.assertIsInstance(result, str)
                followup.assert_not_awaited()
        self.assertEqual(self.mocks["_search_population"].await_count, 5)
        self.mocks["fetch_candidate_ids_by_product"].assert_not_called()

    def test_empty_technical_record_is_not_a_selection(self):
        reports._empty("technical", self.keys[0])
        self.assertTrue(sessions.get_session_extract(self.keys[0]))
        self.assert_format_error()

    def test_specific_id_is_followup(self):
        saved = self.selection()
        prompt = "Расскажи подробнее об обращении 123456"
        self.assertEqual(router.route_intent(prompt, True), router.Intent.SPECIFIC_APPEAL)
        with patch.object(reports, "_run_follow_up", new=AsyncMock(return_value="details")) as followup:
            self.assertEqual(self.run_report(prompt), "details")
            followup.assert_awaited_once_with(self.keys[0], prompt, saved)

    def test_plain_text_not_repaired_into_structured_request(self):
        prompt = "Продукт банковская карта, найди жалобы"
        self.assertFalse(parser.is_structured_analytical_request(prompt))
        with patch.object(reports, "parse_structured_analytical_request", wraps=parser.parse_structured_analytical_request) as parse:
            self.run_report(prompt)
        parse.assert_called_once_with(prompt)
        self.mocks["_search_population"].assert_awaited_once_with(
            self.keys[0], prompt, None, None, structural_filters={})
        self.mocks["fetch_candidate_ids_by_product"].assert_not_called()

    def test_sessions_do_not_interfere(self):
        saved = self.selection(self.keys[1])
        self.assert_format_error()
        self.assertIs(sessions.get_session_extract(self.keys[1]), saved)
        self.assertIn("Готовый отчёт", self.run_report(CANONICAL))
        self.assertIs(sessions.get_session_extract(self.keys[1]), saved)

    def test_all_formats_pass_unmodified_prompt_to_parser(self):
        for prompt in (CANONICAL, JSON_PROMPT, CSV_PROMPT):
            with self.subTest(prompt=prompt), patch.object(reports, "parse_structured_analytical_request", wraps=parser.parse_structured_analytical_request) as parse:
                self.assertIn("Готовый отчёт", self.run_report(prompt))
                parse.assert_called_once_with(prompt)

    def test_invalid_markers_are_new_even_with_selection(self):
        prompts = (INVALID, INVALID_JSON, '{"request_type":"appeals_analysis",',
                   "Анализ обращений\n\nПродукт:\n- Карта", '"bad product", "", "", "query"')
        for prompt in prompts:
            with self.subTest(prompt=prompt):
                self.assertEqual(router.route_intent(prompt, True), router.Intent.NEW_SEARCH)

    def test_filters_dates_and_semantic_query_contract_unchanged(self):
        query = "Первая строка\nВторая строка за 2023 год"
        canonical = "Анализ обращений\n\nПродукт:\n- unknown frontend value\n- Вклады\n\nКанал:\n- Чат\n- Офис\n\nПериод:\nс 01.01.2026\n\nЗапрос:\n" + query
        payload = json.loads(JSON_PROMPT)
        payload["filters"].update(prd=["unknown frontend value", "Вклады"], s_prd=[], chnl=["Чат", "Офис"], date_from="2026-01-01")
        payload["prompt"] = query
        for prompt in (canonical, json.dumps(payload, ensure_ascii=False)):
            request = parser.parse_structured_analytical_request(prompt)
            self.assertEqual(request["query"], query)
            self.assertEqual(request["products"], ["unknown frontend value", "Вклады"])
            self.assertEqual(request["channels"], ["Чат", "Офис"])
            self.assertEqual(request["date_range"], ("2026-01-01", None))
        for prompt in ("Анализ обращений\n\nЗапрос:\n" + query,
                       '"", "", "", "' + query.replace("\n", " ") + '"'):
            self.assertEqual(parser.parse_structured_analytical_request(prompt)["products"], [])

    def test_actual_followup_paths_preserve_saved_ids(self):
        self.selection()
        with patch.object(reports, "classify_complaint_intent", return_value="search"), patch.object(reports, "answer_complaint_follow_up", return_value="found"):
            self.assertEqual(self.run_report("Покажи блокировки в этих обращениях"), "found")
        self.mocks["_search_population"].assert_awaited_once_with(self.keys[0], "Покажи блокировки в этих обращениях", ["123456"], None, structural_filters={})
        with patch.object(reports, "fetch_appeals_by_ids", return_value=self.frame) as fetch, patch.object(reports, "_rows_for_answer", return_value=[]), patch.object(reports, "answer_complaint_details", return_value="details"):
            self.assertEqual(self.run_report("Обращение 123456"), "details")
            fetch.assert_called_once_with(["123456"], None)

    def test_natural_query_has_no_invented_filters_or_mass_population(self):
        for prompt in ("Найди обращения, где клиенты жалуются на мошенничество",
                       "Проанализируй жалобы клиентов на переводы",
                       "Найди обращения по кредитным картам с жалобами на комиссии"):
            with self.subTest(prompt=prompt):
                request = parser.parse_structured_analytical_request(prompt)
                self.assertEqual(request, {"format": "natural", "query": prompt,
                    "products": [], "subproducts": [], "channels": [], "date_range": None})
                sessions.clear_session_extract(self.keys[0])
                self.assertIn("Готовый отчёт", self.run_report(prompt))
                self.mocks["_search_population"].assert_awaited_with(
                    self.keys[0], prompt, None, None, structural_filters={})
        self.mocks["fetch_candidate_ids_by_product"].assert_not_called()
        self.mocks["export_complaints_excel"].assert_called()
        self.mocks["generate_complaint_hypothesis_narrative"].assert_awaited()

    def test_natural_date_is_applied_only_to_candidate_hydration(self):
        query = "Найди обращения о мошенничестве за 2026 год"
        dates = ("2026-01-01", "2026-12-31")
        self.mocks["extract_search_params"].return_value = {"date_range": dates}
        # Exercise the real search function, including None at the OSIRIS boundary.
        source = importlib.import_module(f"{PACKAGE}.utils.srb_d3")
        with patch.object(reports, "_search_population", self.real_search), patch.object(
            reports, "retrieve_via_srb_d3", wraps=source.retrieve_via_srb_d3,
        ) as retrieve, patch.object(source, "ensure_srb_d3_ready"), patch.object(
            source, "submit_request", return_value="request",
        ) as submit, patch.object(source, "wait_result", return_value=["123456"]), patch.object(
            reports, "fetch_appeals_by_ids", return_value=self.frame,
        ) as hydrate, patch.object(reports, "rerank_via_srb_d3", return_value=self.frame):
            self.assertIn("Готовый отчёт", self.run_report(query))
        retrieve.assert_called_once_with(self.keys[0], query, None)
        self.assertIsNone(submit.call_args.args[3]["allowed_ids"])
        hydrate.assert_called_once_with(["123456"], dates)
        self.mocks["fetch_candidate_ids_by_product"].assert_not_called()

    def test_explicit_new_natural_search_overrides_previous_ids(self):
        self.selection()
        prompt = "Сделай новый поиск по всем обращениям о мошенничестве"
        with patch.object(reports, "_run_follow_up", new=AsyncMock()) as followup:
            self.assertIn("Готовый отчёт", self.run_report(prompt))
            followup.assert_not_awaited()
        self.mocks["_search_population"].assert_awaited_once_with(
            self.keys[0], prompt, None, None, structural_filters={})

    def test_among_these_appeals_is_bounded_followup(self):
        self.selection()
        prompt = "А среди этих обращений есть упоминания мошенничества?"
        with patch.object(reports, "answer_complaint_follow_up", return_value="found"):
            self.assertEqual(self.run_report(prompt), "found")
        self.mocks["_search_population"].assert_awaited_once_with(
            self.keys[0], prompt, ["123456"], None, structural_filters={})

    def test_damaged_structured_requests_never_expand_to_global_search(self):
        for prompt in ("Анализ обращений\nНайди жалобы", '{"request_type":"appeals_analysis",',
                       '"Кредиты", "", "найди жалобы"', "Продукт:\n- Кредиты\nЗапрос:\nНайди жалобы",
                       "Кредиты, , , Найди жалобы"):
            with self.subTest(prompt=prompt):
                self.selection()
                self.assert_format_error(prompt)
                self.assertFalse(sessions.get_session_extract(self.keys[0]).get("final_ids"))

    def test_natural_osiris_outage_is_not_a_fictional_report(self):
        from workspace.utils.osiris_runtime import OsirisUnavailableError
        self.mocks["_search_population"].side_effect = OsirisUnavailableError("offline")
        result = self.run_report("Найди обращения о мошенничестве")
        self.assertIn("сервис обработки недоступен", result)
        self.mocks["export_complaints_excel"].assert_not_called()
        self.mocks["generate_complaint_hypothesis_narrative"].assert_not_awaited()
        self.assertFalse(sessions.get_session_extract(self.keys[0]).get("final_ids"))

    def test_tool_accepts_any_nonempty_text_after_agent_routing(self):
        for text in ("Какая сегодня погода?", "2026", "Найди документацию Python"):
            with self.subTest(text=text):
                self.assertEqual(parser.parse_structured_analytical_request(text)["query"], text)
        for text in ("", "  \n"):
            with self.assertRaises(ValueError):
                parser.parse_structured_analytical_request(text)

    def test_absent_period_never_asks_llm_to_invent_dates(self):
        self.assertEqual(self.real_extract_dates("Найди обращения о мошенничестве"), {"date_range": None})
        self.mocks["def_ask_gigachat"].assert_not_called()

    def test_natural_search_produces_normal_xlsx_and_hypotheses(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory, patch.object(
            reports, "export_complaints_excel", self.real_export,
        ), patch.object(reports, "output_directory", return_value=Path(directory)), patch.object(
            reports, "register_artifact",
        ) as register:
            result = self.run_report("Проанализируй жалобы клиентов на переводы")
            self.assertIn("Готовый отчёт", result)
            register.assert_called_once()
            workbook = load_workbook(register.call_args.args[0])
            self.assertEqual(workbook.active.max_row, 2)
            workbook.close()
        self.mocks["generate_complaint_hypothesis_narrative"].assert_awaited_once()

    def test_router_is_stdlib_only_and_never_invokes_parser(self):
        source = (SKILL / "utils/intent_router.py").read_text(encoding="utf-8")
        tree = ast.parse(source)
        imports = {node.module for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)}
        imports.update(alias.name for node in ast.walk(tree) if isinstance(node, ast.Import) for alias in node.names)
        self.assertLessEqual(imports, {"__future__", "enum", "json", "re"})
        with patch.object(parser, "parse_structured_analytical_request", side_effect=AssertionError("Router must not parse")):
            for prompt in (INVALID, CANONICAL, INVALID_JSON, JSON_PROMPT, CSV_PROMPT):
                self.assertEqual(router.route_intent(prompt, True), router.Intent.NEW_SEARCH)
        self.mocks["def_ask_gigachat"].assert_not_called()
        self.mocks["fetch_candidate_ids_by_product"].assert_not_called()

    def test_ordinary_followup_is_not_a_new_envelope(self):
        self.assertFalse(router.looks_like_new_analytical_request("Найди жалобы на карты"))
        self.assertEqual(router.route_intent("Найди жалобы на карты", True), router.Intent.NEW_SEARCH)
        for prompt in ("Какие основные причины жалоб среди них?",
                       "Расскажи подробнее об обращении 123456", "Что означает Анализ обращений?"):
            self.assertFalse(router.looks_like_new_analytical_request(prompt))
            self.assertNotEqual(router.route_intent(prompt, True), router.Intent.NEW_SEARCH)

    def test_new_topic_without_reference_is_global_and_natural_variants_are_accepted(self):
        for prompt in ("Найди обращения с жалобами на ипотеку", "Есть ли обращения с жалобами на блокировку карт?",
                       "Посмотри обращения по мошенничеству", "Проверь жалобы клиентов на комиссии",
                       "Какие обращения связаны с переводами?", "Нужна аналитика обращений о проблемах с кредитами",
                       "Найди обращения о комиссиях, переводах, блокировках и мошенничестве",
                       "Найди обращения с жалобами на ипотеку свыше 100000 рублей"):
            with self.subTest(prompt=prompt):
                self.selection()
                self.assertEqual(router.route_intent(prompt, True), router.Intent.NEW_SEARCH)
                self.assertEqual(parser.parse_structured_analytical_request(prompt)["format"], "natural")
                self.assertIn("Готовый отчёт", self.run_report(prompt))
                self.mocks["_search_population"].assert_awaited_with(self.keys[0], prompt, None, None, structural_filters={})
        self.mocks["fetch_candidate_ids_by_product"].assert_not_called()

    def test_found_earlier_reference_is_followup_and_date_empty_result_is_qualified(self):
        self.assertEqual(router.route_intent("А среди найденных ранее есть мошенничество?", True), router.Intent.FOLLOWUP_SEARCH)
        self.mocks["extract_search_params"].return_value = {"date_range": ("2026-01-01", "2026-12-31")}
        self.mocks["_search_population"].return_value = pd.DataFrame()
        result = self.run_report("Найди обращения по мошенничеству за 2026 год")
        self.assertIn("не доказывает отсутствие", result)

    def test_new_all_formats_always_override_saved_selection(self):
        self.selection()
        for prompt in (CANONICAL, JSON_PROMPT, CSV_PROMPT):
            self.assertEqual(router.route_intent(prompt, True), router.Intent.NEW_SEARCH)

    def test_native_tool_preserves_prompt_and_session_on_retries(self):
        base = ModuleType("nanobot.agent.tools.base")
        base.Tool = type("Tool", (), {})
        base.tool_parameters = lambda schema: lambda cls: cls
        base.ToolResult = SimpleNamespace(error=lambda message: SimpleNamespace(is_error=True, message=message))
        context = ModuleType("nanobot.agent.tools.context")
        context.current_request_context = lambda: object()
        context.current_request_session_key = lambda: self.keys[0]
        spec = importlib.util.spec_from_file_location("appeals_native_routing_test", SKILL.parents[1] / "tools/appeals_analyzer.py")
        native = importlib.util.module_from_spec(spec)
        with patch.dict(sys.modules, {base.__name__: base, context.__name__: context}):
            spec.loader.exec_module(native)
        artifacts = importlib.import_module(f"{PACKAGE}.utils.appeals_artifacts")
        backend = importlib.import_module(f"{PACKAGE}.utils.data_store")
        tool = native.AppealsAnalyzerTool(config=native.AppealsAnalyzerToolConfig(), message_tool=SimpleNamespace(execute=AsyncMock()))
        runner = AsyncMock(side_effect=["Ошибка формата", "Ошибка формата", "Отчёт"])
        with patch.object(native, "log_skill_runtime", return_value="production"), patch.object(tool, "_load_runner", return_value=runner), patch.object(native, "import_module", side_effect=lambda name: artifacts if name.endswith("appeals_artifacts") else backend):
            for prompt in (INVALID, INVALID, CANONICAL):
                asyncio.run(tool.execute(prompt=prompt, session_id="ignored-agent-id"))
        self.assertEqual([call.kwargs for call in runner.await_args_list], [
            {"session_id": self.keys[0], "user_prompt": prompt} for prompt in (INVALID, INVALID, CANONICAL)])
        tool._message_tool.execute.assert_not_awaited()


if __name__ == "__main__":
    unittest.main()

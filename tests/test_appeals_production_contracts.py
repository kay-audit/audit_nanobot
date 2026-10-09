"""Offline production contracts, runnable with unittest without Gateway/ML packages."""
from __future__ import annotations

import ast
import asyncio
import importlib
import importlib.util
import json
import pickle
import subprocess
import sys
import tempfile
import time
import types
import unittest
import zipfile
from dataclasses import replace
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import AsyncMock, Mock, patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
SKILL = ROOT / "workspace/skills/appeals-analyzer"
PACKAGE = "appeals_production_contract"


def load_package():
    if PACKAGE not in sys.modules:
        spec = importlib.util.spec_from_file_location(
            PACKAGE, SKILL / "__init__.py", submodule_search_locations=[str(SKILL)],
        )
        module = importlib.util.module_from_spec(spec)
        sys.modules[PACKAGE] = module
        spec.loader.exec_module(module)


def load_file(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


load_package()
sql = importlib.import_module(f"{PACKAGE}.utils.greenplum_engine")
store = importlib.import_module(f"{PACKAGE}.utils.data_store")
text = importlib.import_module(f"{PACKAGE}.utils.appeal_text")
reports = importlib.import_module(f"{PACKAGE}.scripts.appeals_reports")
sessions = importlib.import_module(f"{PACKAGE}.utils.session_extract_manager")
artifacts = importlib.import_module(f"{PACKAGE}.utils.appeals_artifacts")
client = importlib.import_module(f"{PACKAGE}.utils.srb_d3")
bge = importlib.import_module(f"{PACKAGE}.utils.bge_search_engine")
worker = load_file("appeals_worker_contract", SKILL / "rerank_osiris_worker.py")

PROMPT = """Анализ обращений

Продукт:
- Кредиты
- Вклады

Подпродукт:
- Дебетовая карта

Канал:
- Чат
- Офис

Период:
01.01.2026 — 31.07.2026

Запрос:
жалобы за 2023 год
"""


class ParserAndSQLTests(unittest.TestCase):
    def test_authoritative_canonical_values_and_period(self):
        request = sql.parse_structured_analytical_request(PROMPT)
        self.assertEqual(request["channels"], ["Чат", "Офис"])
        self.assertEqual(request["products"], ["Кредиты", "Вклады"])
        self.assertEqual(request["subproducts"], ["Дебетовая карта"])
        self.assertEqual(request["date_range"], ("2026-01-01", "2026-07-31"))

    def test_groups_and_dates_and_text_population(self):
        query, params = sql.build_product_prefilter_sql(
            ["Кредиты", "Вклады"], ["Дебетовая карта"], ["Чат", "Офис"],
            date_range=("2026-01-01", "2026-07-31"),
        )
        self.assertIn("a.prd IN (%s, %s) AND a.s_prd IN (%s) AND a.chnl IN (%s, %s)", query)
        self.assertIn("a.app_row_id AS VARCHAR", query)
        self.assertNotIn("a.id", query)
        self.assertIn("EXISTS", query)
        self.assertIn("msg_crm_call", query)
        self.assertEqual(str(params[-1]), "2026-08-01")
        self.assertNotIn("appeal_2025", query)

    def test_date_only_and_empty_groups(self):
        for period in (None, ("2026-01-01", None), (None, "2026-07-31")):
            query, _ = sql.build_product_prefilter_sql([], [], [], date_range=period)
            self.assertIn("EXISTS", query)
            self.assertNotIn("prd IN", query)

    def test_legacy_still_uses_dictionary(self):
        with self.assertRaises(ValueError):
            sql.parse_structured_analytical_request('"несуществующий продукт", "", "", "query"')

    def test_hydration_uses_only_canonical_id_and_required_columns(self):
        query = sql.build_hydration_sql(["123", "O'Reilly"])
        self.assertIn("a.app_row_id AS VARCHAR) AS id", query)
        self.assertIn("O''Reilly", query)
        self.assertNotIn("a.id", query)
        self.assertIn("a.kanal_reg", query)
        self.assertIn("a.req_desc", query)
        tasks = sql.build_task_hydration_sql({2026: ["123"]})
        self.assertIn("task_answer_full", tasks)
        self.assertNotIn("task_text_sol", tasks)

    def test_relation_aggregation_and_crm_fallback(self):
        base = pd.DataFrame([{"source_year": 2026, "_join_app_row_id": "123", "app_row_id": "123",
                              "id": "123", "req_desc": "описание"}])
        dialogs = pd.DataFrame([
            {"source_year": 2026, "_join_app_row_id": "123", "msg_pprb_chat": None, "msg_crm_call": "звонок"},
            {"source_year": 2026, "_join_app_row_id": "123", "msg_pprb_chat": "", "msg_crm_call": "звонок"},
        ])
        tasks = pd.DataFrame([
            {"source_year": 2026, "_join_app_row_id": "123", "task_answer": "a", "task_answer_full": "A"},
            {"source_year": 2026, "_join_app_row_id": "123", "task_answer": "b", "task_answer_full": "B"},
        ])
        result = sql.merge_hydration_frames(base, dialogs, tasks)
        self.assertEqual(len(result), 1)
        self.assertEqual(result.iloc[0]["app_row_id"], "123")
        self.assertEqual(result.iloc[0]["Транскрибация диалога"], "звонок")
        self.assertEqual(len(result.iloc[0]["tasks"]), 2)

    def test_text_helper(self):
        for chat in (None, np.nan, pd.NA, "", "  "):
            self.assertEqual(text.canonical_appeal_text({"req_desc": " desc ", "msg_pprb_chat": chat,
                                                       "msg_crm_call": "crm"}), "desc crm")
        self.assertEqual(text.canonical_appeal_text({"req_desc": "desc", "msg_pprb_chat": "chat",
                                                   "msg_crm_call": "crm"}), "desc chat")
        self.assertEqual(text.canonical_appeal_text({"req_desc": np.nan, "msg_pprb_chat": None}), "")

    def test_unavailable_year_skips_queries(self):
        with store.backend_scope("greenplum"), patch.object(sql, "configured_years", return_value=[2026]), patch.object(sql, "_run_sql") as query:
            self.assertEqual(sql.fetch_candidate_ids_by_product([], [], [], ("2025-01-01", "2025-12-31")), [])
            self.assertTrue(sql.fetch_appeals_by_ids(["1"], ("2025-01-01", "2025-12-31")).empty)
        query.assert_not_called()


class BackendTests(unittest.TestCase):
    def test_provider_wait_and_dataframe_contract(self):
        provider = Mock()
        provider.open_cache.side_effect = [False, True]
        columns = ("app_row_id", "req_reg_date", "prd", "s_prd", "chnl")
        types = ("VARCHAR", "TIMESTAMP", "VARCHAR", "VARCHAR", "VARCHAR")
        provider.query_sql.side_effect = [
            {"status": "success", "columns": ["column_name", "data_type"],
             "rows": [{"column_name": name, "data_type": kind} for name, kind in zip(columns, types)]},
            {"status": "success", "rows": [{"id": "123"}], "columns": ["id"]},
        ]
        with patch.object(store, "source_years", return_value=[2026]):
            cache = store.SharedCacheStore(provider, wait_seconds=1, poll_interval=.001)
        frame = cache.query_sql("SELECT id WHERE id=%s", ["123"])
        self.assertEqual(frame.id.tolist(), ["123"])
        provider.query_sql.assert_called_with("SELECT id WHERE id=%s", ["123"])
        provider.refresh.assert_not_called()

    def test_missing_snapshot_fails_without_refresh(self):
        provider = Mock()
        provider.open_cache.return_value = False
        with self.assertRaisesRegex(RuntimeError, "snapshot is unavailable"):
            store.SharedCacheStore(provider, wait_seconds=0)
        provider.refresh.assert_not_called()

    def test_missing_table_is_an_error(self):
        provider = Mock()
        provider.open_cache.return_value = True
        provider.query_sql.return_value = {"status": "error", "error": "table missing"}
        with patch.object(store, "source_years", return_value=[2026]):
            with self.assertRaisesRegex(RuntimeError, "table missing"):
                store.SharedCacheStore(provider)

    def test_default_cache_error_never_falls_back_to_gp(self):
        forbidden = types.ModuleType(f"{PACKAGE}.utils.db")
        forbidden.run = Mock(side_effect=AssertionError("GP called"))
        with patch.dict(sys.modules, {forbidden.__name__: forbidden}), patch.object(store, "_store", None), \
                patch.object(store, "SharedCacheStore", side_effect=RuntimeError("cache missing")):
            with self.assertRaisesRegex(RuntimeError, "cache missing"):
                store.query_sql("SELECT 1")
        forbidden.run.assert_not_called()

    def test_explicit_gp_backend_and_scope_reset(self):
        cursor = Mock()
        cursor.description = [("id",)]
        cursor.fetchall.return_value = [("123",)]
        cursor.__enter__ = Mock(return_value=cursor)
        cursor.__exit__ = Mock(return_value=False)
        connection = Mock()
        connection.cursor.return_value = cursor
        fake = types.ModuleType(f"{PACKAGE}.utils.db")
        fake.run = lambda fn: fn(connection)
        with patch.dict(sys.modules, {fake.__name__: fake}), store.backend_scope("greenplum"):
            self.assertEqual(store.query_sql("SELECT id", []).id.tolist(), ["123"])
        self.assertEqual(store._backend.get(), "cache")

    def test_backend_context_survives_to_thread_and_isolated_requests(self):
        async def run(name):
            with store.backend_scope(name):
                await asyncio.sleep(.001)
                return await asyncio.to_thread(store._backend.get)

        async def all_calls():
            return await asyncio.gather(run("cache"), run("greenplum"))

        self.assertEqual(asyncio.run(all_calls()), ["cache", "greenplum"])

    def test_hydration_uses_selected_backend_for_all_three_queries(self):
        base = pd.DataFrame([{"source_year": 2026, "_join_app_row_id": "123", "id": "123", "app_row_id": "123"}])
        with patch.object(sql, "configured_years", return_value=[2026]), \
                patch.object(sql, "query_sql", side_effect=[base, pd.DataFrame(), pd.DataFrame()]) as query:
            result = sql.fetch_appeals_by_ids(["123"])
        self.assertEqual(query.call_count, 3)
        self.assertEqual(result.id.tolist(), ["123"])

    def test_empty_hydration_skips_relations(self):
        with patch.object(sql, "configured_years", return_value=[2026]), \
                patch.object(sql, "query_sql", return_value=pd.DataFrame()) as query:
            self.assertTrue(sql.fetch_appeals_by_ids(["123"]).empty)
        self.assertEqual(query.call_count, 1)


class Selector:
    def __init__(self, positions):
        self.positions = list(map(int, positions))


class FaissIndex:
    ntotal = 6
    nprobe = 7
    d = 2

    def __init__(self):
        self.selections = []

    def search(self, vector, k, params=None):
        selected = list(range(self.ntotal)) if params is None else params.sel.positions
        self.selections.append(selected)
        return np.zeros((1, min(k, len(selected)))), np.asarray([selected[:k]], dtype=int)


class BM25Shard:
    def __init__(self):
        self.scores = {"num_docs": 3}
        self.masks = []

    def retrieve(self, query, k, weight_mask=None, **kwargs):
        self.masks.append(weight_mask)
        selected = [i for i in reversed(range(3)) if weight_mask is None or weight_mask[i]]
        return np.asarray([selected[:k]], dtype=int), np.zeros((1, min(k, len(selected))))


class RetrievalTests(unittest.TestCase):
    def setUp(self):
        self.index = FaissIndex()
        self.shards = [BM25Shard(), BM25Shard()]
        self.faiss = types.SimpleNamespace(IDSelectorBatch=Selector, SearchParametersIVF=types.SimpleNamespace)
        self.embed = Mock()
        self.embed.encode.return_value = np.zeros((1, 2), dtype="float32")
        self.patches = [
            patch.dict(sys.modules, {"faiss": self.faiss}),
            patch.object(bge, "doc_ids", ["a", "b", "c", "d", "e", "f"]),
            patch.object(bge, "id_to_positions", {cid: i for i, cid in enumerate("abcdef")}),
            patch.object(bge, "faiss_loaded", self.index),
            patch.object(bge, "bm25_indexes", [(self.shards[0], 0), (self.shards[1], 3)]),
            patch.object(bge, "_BGE_CACHE", {"embed": self.embed}),
        ]
        for p in self.patches:
            p.start()
            self.addCleanup(p.stop)

    def test_selector_shard_masks_and_unknown_id(self):
        result = bge.retrieve_hybrid_adaptive("query", ["a", "c", "e", "missing"])
        self.assertEqual(set(result), {"a", "c", "e"})
        self.assertEqual(self.index.selections, [[0, 2, 4]])
        self.assertEqual(self.shards[0].masks[0].tolist(), [1, 0, 1])
        self.assertEqual(self.shards[1].masks[0].tolist(), [0, 1, 0])

    def test_empty_is_not_unrestricted(self):
        for ids in ([], ["missing"]):
            self.assertEqual(bge.retrieve_hybrid_adaptive("query", ids), [])
        self.embed.encode.assert_not_called()
        self.assertEqual(self.index.selections, [])

    def test_unrestricted_and_no_candidate_cap(self):
        self.assertEqual(len(bge.retrieve_hybrid_adaptive("query", None)), 6)
        self.assertIsNone(self.shards[0].masks[0])
        ranks = {str(i): i + 1 for i in range(2048)}
        self.assertEqual(len(bge.fuse_rrf_rank_maps(ranks, {})), 2048)

    def test_reference_parameters(self):
        self.assertEqual((bge.CONFIG.faiss_k, bge.CONFIG.bm25_total_k, bge.CONFIG.rrf_k,
                          bge.CONFIG.rrf_alpha, bge.CONFIG.score_threshold), (2048, 1372, 60, .3, .5))

    def test_algorithm_matches_reference_notebook(self):
        archive_path = ROOT / "workspace/reference/faiss_bm25_d3.zip"
        if not archive_path.is_file():
            self.skipTest("Local reference archive is not distributed with the repository")
        with zipfile.ZipFile(archive_path) as archive:
            notebook = json.loads(archive.read("faiss_bm25_d3/pipeline_search.ipynb"))
        source = next("".join(cell["source"]) for cell in notebook["cells"]
                      if "def retrieve_hybrid_adaptive(" in "".join(cell.get("source", [])))
        function = next(node for node in ast.parse(source).body
                        if isinstance(node, ast.FunctionDef) and node.name == "retrieve_hybrid_adaptive")
        env = {"np": np, "doc_ids": bge.doc_ids, "faiss": self.faiss, "ALPHA": .3, "K_RRF": 60,
               "embed": lambda texts: np.zeros((1, 2), dtype="float32"),
               "tokenize": bge.tokenize, "build_date_mask": lambda period: np.array([1, 0, 1, 0, 1, 0], dtype=bool)}
        exec(compile(ast.Module(body=[function], type_ignores=[]), "reference", "exec"), env)
        expected = env["retrieve_hybrid_adaptive"]("query", self.index, bge.bm25_indexes, 99999, ("start", "end"))
        self.assertEqual(bge.retrieve_hybrid_adaptive("query", ["a", "c", "e"]), expected)

    def test_gpu_without_selector_support_uses_cpu_inside_worker(self):
        self.faiss.StandardGpuResources = lambda: object()
        self.faiss.index_cpu_to_gpu = Mock(side_effect=RuntimeError("GPU unsupported"))
        with patch.object(bge, "load_pipeline_meta_and_indices"), patch.object(bge, "_gpu_resources", None):
            bge.initialize_retrieval(self.embed)
        self.assertIs(bge.faiss_loaded, self.index)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        sessions.clear_session_extract("contract")

    def test_minimum_500_keeps_low_scores_and_orders_by_score(self):
        frame = pd.DataFrame({"id": ["a", "b", "c"], "score": [.49, .5, .9]})
        self.assertEqual(reports.select_accepted(frame).id.tolist(), ["c", "b", "a"])
        self.assertEqual(reports.select_accepted(frame.iloc[:1]).id.tolist(), ["a"])
        self.assertTrue(reports.select_accepted(frame.iloc[:0]).empty)
        large = pd.DataFrame({"id": range(700), "score": np.linspace(0, .49, 700)})
        selected = reports.select_accepted(large)
        self.assertEqual(len(selected), 500)
        self.assertEqual(selected.id.tolist(), list(range(699, 199, -1)))

    def test_minimum_500_deduplicates_before_counting_and_keeps_best_duplicate(self):
        frame = pd.DataFrame({"id": ["duplicate"] * 600 + list(range(600)),
                              "score": [.9] * 600 + [.01] * 600})
        selected = reports.select_accepted(frame)
        self.assertEqual(len(selected), 500)
        self.assertTrue(selected.id.is_unique)
        self.assertEqual(selected.id.tolist(), ["duplicate"] + list(range(499)))
        frame = pd.DataFrame({"id": ["a", "a", "b"], "score": [.01, .2, 0]})
        self.assertEqual(reports.select_accepted(frame).score.tolist(), [.2, 0])

    def test_all_threshold_passes_are_retained_above_500_without_cap(self):
        for passing_count in (0, 1, 499, 500, 501, 700):
            with self.subTest(passing_count=passing_count):
                frame = pd.DataFrame({"id": range(passing_count + 700),
                                      "score": [.500001] * passing_count + [.01] * 700})
                selected = reports.select_accepted(frame)
                self.assertEqual(len(selected), max(500, passing_count))
                self.assertEqual(int((selected.score > .5).sum()), passing_count)
                self.assertTrue(selected.id.is_unique)

    def test_report_selection_required_populations(self):
        cases = [(750, 750, 750), (1000, 700, 700), (1000, 13, 500),
                 (1000, 0, 500), (320, 10, 320), (4347, 13, 500), (4347, 1200, 1200)]
        for total, above, expected in cases:
            with self.subTest(total=total, above=above):
                frame = pd.DataFrame({"id": [str(i) for i in range(total)],
                                      "score": [.9] * above + np.linspace(0, .5, total - above).tolist()})
                frame = frame.sample(frac=1, random_state=42).reset_index(drop=True)
                original = frame.copy(deep=True)
                result = reports.select_accepted(frame)
                self.assertEqual(len(result), expected)
                self.assertTrue(result.id.is_unique)
                self.assertTrue(result.score.is_monotonic_decreasing)
                high = frame[frame.score > .5]
                self.assertTrue(set(high.id).issubset(set(result.id)))
                if above < 500:
                    expected_ids = frame.sort_values("score", ascending=False, kind="stable").head(expected).id.tolist()
                    self.assertEqual(result.id.tolist(), expected_ids)
                else:
                    self.assertEqual(set(result.id), set(high.id))
                pd.testing.assert_frame_equal(frame, original)

    def test_report_threshold_is_strict_at_half(self):
        frame = pd.DataFrame({"id": [str(i) for i in range(502)],
                              "score": [.9] * 499 + [.500001, .5, .499999]})
        result = reports.select_accepted(frame)
        self.assertEqual(len(result), 500)
        self.assertIn("499", result.id.tolist())
        self.assertNotIn("500", result.id.tolist())
        self.assertNotIn("501", result.id.tolist())

    def test_invalid_scores_are_errors(self):
        for value in (np.nan, float("inf"), -1, 2):
            with self.assertRaises(RuntimeError):
                reports.select_accepted(pd.DataFrame({"id": ["1"], "score": [value]}))

    def test_startup_unavailable_returns_retry_message_without_export(self):
        with patch.object(reports, "fetch_candidate_ids_by_product", return_value=["123"]), \
                patch.object(reports, "retrieve_via_srb_d3", side_effect=client.OsirisUnavailableError("SDK down")), \
                patch.object(reports, "export_complaints_excel") as export:
            result = asyncio.run(reports.run_appeals_report("contract", PROMPT))
        self.assertIn("через 15 минут", result)
        export.assert_not_called()

    def test_ready_worker_handler_error_is_not_disguised_as_startup_failure(self):
        with patch.object(reports, "fetch_candidate_ids_by_product", return_value=["123"]), \
                patch.object(reports, "retrieve_via_srb_d3", side_effect=client.OsirisRequestError("index broken")):
            with self.assertRaisesRegex(client.OsirisRequestError, "index broken"):
                asyncio.run(reports.run_appeals_report("contract", PROMPT))

    def test_canonical_date_no_llm_and_low_scores_still_export(self):
        frame = pd.DataFrame({"id": ["123"], "app_row_id": ["123"], "req_desc": ["text"]})
        with patch.object(reports, "extract_search_params", side_effect=AssertionError("LLM dates")), \
                patch.object(reports, "fetch_candidate_ids_by_product", return_value=["123"]) as prefilter, \
                patch.object(reports, "retrieve_via_srb_d3", return_value=["123"]) as retrieve, \
                patch.object(reports, "fetch_appeals_by_ids", return_value=frame) as hydrate, \
                patch.object(reports, "rerank_via_srb_d3", return_value=frame.assign(score=.49)), \
                patch.object(reports, "export_complaints_excel", return_value={"name": "appeals.xlsx", "count": 1}) as export, \
                patch.object(reports, "generate_complaint_hypothesis_narrative", new=AsyncMock(return_value="hypotheses")) as narrative:
            result = asyncio.run(reports.run_appeals_report("contract", PROMPT))
        self.assertIn("appeals.xlsx (1 уникальных обращений)", result)
        self.assertEqual(prefilter.call_args.args[-1], ("2026-01-01", "2026-07-31"))
        self.assertEqual(hydrate.call_args.args[-1], ("2026-01-01", "2026-07-31"))
        self.assertEqual(retrieve.call_args.args, ("contract", "жалобы за 2023 год", ["123"]))
        export.assert_called_once()
        self.assertEqual(export.call_args.args[0].score.tolist(), [.49])
        narrative.assert_awaited_once()

    def test_xlsx_and_final_session_ids(self):
        frame = pd.DataFrame({"id": ["123", "456"], "app_row_id": ["123", "456"],
                              "req_desc": ["desc", "desc"], "score": [.9, .49]})
        with tempfile.TemporaryDirectory() as directory:
            paths = []
            with artifacts.artifact_scope(Path(directory), paths), \
                    patch.object(reports, "fetch_candidate_ids_by_product", return_value=["123", "456"]), \
                    patch.object(reports, "retrieve_via_srb_d3", return_value=["123", "456"]), \
                    patch.object(reports, "fetch_appeals_by_ids", return_value=frame), \
                    patch.object(reports, "rerank_via_srb_d3", return_value=frame), \
                    patch.object(reports, "generate_complaint_hypothesis_narrative", new=AsyncMock(return_value="4 hypotheses")):
                report = asyncio.run(reports.run_appeals_report("contract", PROMPT))
            self.assertEqual(len(paths), 1)
            self.assertEqual(list(Path(directory).glob("*.csv")), [])
            workbook = load_workbook(paths[0])
            self.assertEqual(workbook.active.max_row, 3)
            workbook.close()
            self.assertNotIn(directory, report)
            self.assertEqual(sessions.get_session_extract("contract")["final_ids"], ["123", "456"])

    def test_followup_uses_final_ids_only(self):
        sessions.set_session_extract("contract", pd.DataFrame(), extra={"final_ids": ["123"], "hypothesis": "h"})
        with patch.object(reports, "fetch_candidate_ids_by_product", side_effect=AssertionError("new population")), \
                patch.object(reports, "retrieve_via_srb_d3", return_value=[]) as retrieve:
            result = asyncio.run(reports.run_appeals_report("contract", "найди среди них новую тему"))
        self.assertEqual(retrieve.call_args.args[-1], ["123"])
        self.assertEqual(result, "Релевантные обращения не подтверждены.")

    def test_id_lookup_is_exact(self):
        sessions.set_session_extract("contract", pd.DataFrame(), extra={"final_ids": ["123"], "hypothesis": "h"})
        with patch.object(reports, "fetch_appeals_by_ids") as hydrate, \
                patch.object(reports, "answer_complaint_dialog", return_value="dialog"):
            self.assertEqual(asyncio.run(reports.run_appeals_report("contract", "обращение 1234")), "dialog")
        hydrate.assert_not_called()

    def test_import_has_no_external_bootstrap_or_ml(self):
        code = f"""import importlib.util,importlib,sys,socket,threading
def forbidden(*a,**k): raise AssertionError('external side effect')
threading.Thread.start=forbidden
socket.create_connection=forbidden
p={str(SKILL)!r}
s=importlib.util.spec_from_file_location('safe_appeals',p+'/__init__.py',submodule_search_locations=[p])
m=importlib.util.module_from_spec(s);sys.modules[s.name]=m;s.loader.exec_module(m)
importlib.import_module('safe_appeals.scripts.appeals_reports')
assert not any(n in sys.modules for n in ['torch','faiss','bm25s','sentence_transformers'])
assert 'safe_appeals.utils.db' not in sys.modules
"""
        result = subprocess.run([sys.executable, "-B", "-c", code], capture_output=True, text=True, cwd=ROOT)
        self.assertEqual(result.returncode, 0, result.stderr)


class ProtocolTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        profile = replace(client.SERVICE, nfs_root=self.root)
        for module, values in (
            (client, {"SERVICE": profile}),
            (worker, {"PROFILE": profile, "NFS_ROOT": self.root,
                      "INBOX": self.root / "inbox", "PROCESSING": self.root / "processing"}),
        ):
            for name, value in values.items():
                p = patch.object(module, name, value)
                p.start()
                self.addCleanup(p.stop)

    def test_retrieve_request_roundtrip(self):
        request_id = client.submit_request("user:a", "retrieve", "query", {"allowed_ids": ["1", "2"]})
        search = Mock()
        search.retrieve_hybrid_adaptive.return_value = ["2"]
        worker.process_request(self.root / "inbox" / f"{request_id}.json", search, None)
        self.assertEqual(client.wait_result("user:a", request_id, "retrieve"), ["2"])
        self.assertEqual(list(self.root.rglob("*.pkl")), [])
        search.retrieve_hybrid_adaptive.assert_called_once_with("query", ["1", "2"])

    def test_reranker_score_roundtrip_and_order(self):
        request_id = client.submit_request("a", "rerank", "query", {"items": [
            {"id": "123", "text": "text"}, {"id": "456", "text": "text"},
        ]})
        with patch.object(worker, "predict_with_retry", return_value=np.array([0, -1])):
            worker.process_request(self.root / "inbox" / f"{request_id}.json", None, object())
        result = client.wait_result("a", request_id, "rerank")
        self.assertEqual(result.id.tolist(), ["123", "456"])
        self.assertEqual(result.score.iloc[0], .5)
        self.assertLess(result.score.iloc[1], .5)

    def test_worker_and_adapter_return_all_candidates_including_low_scores(self):
        count = 4347
        probabilities = np.resize(np.array([.9, .5, .4, .2, .01]), count)
        logits = np.log(probabilities / (1 - probabilities))
        frame = pd.DataFrame({"id": [str(i) for i in range(count)], "req_desc": ["text"] * count})
        original_wait = client.wait_result

        def score_and_wait(session_id, request_id, request_type, *args, **kwargs):
            worker.process_request(self.root / "inbox" / f"{request_id}.json", None, object())
            return original_wait(session_id, request_id, request_type, *args, **kwargs)

        with patch.object(client, "ensure_srb_d3_ready"), \
                patch.object(client, "wait_result", side_effect=score_and_wait), \
                patch.object(worker, "predict_with_retry", return_value=logits) as score:
            result = client.rerank_via_srb_d3("low-score-contract", "query", frame)
        self.assertEqual(len(score.call_args.args[1]), count)
        self.assertEqual(len(result), count)
        self.assertEqual(set(result.id), set(frame.id))
        self.assertTrue(result.score.is_monotonic_decreasing)
        correlated = result.set_index("id").loc[frame.id, "score"].to_numpy()
        np.testing.assert_allclose(correlated, probabilities)
        for value in (.4, .2, .01):
            self.assertTrue(np.isclose(result.score, value).any())

    def test_client_hydration_scores_join_by_id(self):
        frame = pd.DataFrame({"id": ["old1", "old2"], "app_row_id": ["123", "456"],
                              "req_desc": ["desc", "other"], "msg_pprb_chat": [None, "chat"],
                              "msg_crm_call": ["crm", "unused"]})
        with patch.object(client, "ensure_srb_d3_ready"), patch.object(client, "submit_request", return_value="r") as submit, \
                patch.object(client, "wait_result", return_value=pd.DataFrame({"id": ["456", "123"], "score": [.9, .5]})):
            result = client.rerank_via_srb_d3("a", "query", frame)
        self.assertEqual(result.id.tolist(), ["456", "123"])
        self.assertEqual(submit.call_args.args[3]["items"][0], {"id": "123", "text": "desc crm"})

    def test_response_correlation_rejected(self):
        request_id = client.submit_request("a", "retrieve", "q", {"allowed_ids": None})
        _, dirs = client._session_dirs("a")
        client._atomic_pickle({"request_id": "wrong", "request_type": "retrieve", "result": []},
                              dirs["output"] / f"{request_id}.pkl")
        with self.assertRaisesRegex(RuntimeError, "correlation"):
            client.wait_result("a", request_id, "retrieve")

    def test_timeout_cleans_queued_request(self):
        request_id = client.submit_request("a", "retrieve", "q", {"allowed_ids": None}, .001)
        with self.assertRaises(TimeoutError):
            client.wait_result("a", request_id, "retrieve", .001)
        self.assertEqual(list(self.root.rglob("*.pkl")), [])
        self.assertEqual(list((self.root / "inbox").glob("*.json")), [])

    def test_expired_request_does_not_execute(self):
        request_id = client.submit_request("a", "retrieve", "q", {"allowed_ids": None}, .001)
        search = Mock()
        time.sleep(.005)
        worker.process_request(self.root / "inbox" / f"{request_id}.json", search, None)
        search.retrieve_hybrid_adaptive.assert_not_called()
        self.assertEqual(list(self.root.rglob("*.pkl")), [])

    def test_worker_errors_are_correlated_and_cleaned(self):
        request_id = client.submit_request("a", "retrieve", "q", {"allowed_ids": None})
        search = Mock()
        search.retrieve_hybrid_adaptive.side_effect = RuntimeError("index broken")
        with patch("builtins.print"):
            worker.process_request(self.root / "inbox" / f"{request_id}.json", search, None)
        with self.assertRaisesRegex(RuntimeError, "request_id=" + request_id):
            client.wait_result("a", request_id, "retrieve")
        self.assertEqual(list(self.root.rglob("*.pkl")), [])
        self.assertEqual(list(self.root.rglob("*.txt")), [])

    def test_parallel_sessions_have_distinct_requests(self):
        with ThreadPoolExecutor(max_workers=4) as executor:
            requests = list(executor.map(lambda session: (session, client.submit_request(
                session, "retrieve", "q", {"allowed_ids": [session]},
            )), ["a", "b", "a", "b"]))
        self.assertEqual(len({request for _, request in requests}), 4)
        search = types.SimpleNamespace(retrieve_hybrid_adaptive=lambda q, ids: ids)
        for session, request in reversed(requests):
            worker.process_request(self.root / "inbox" / f"{request}.json", search, None)
        for session, request in requests:
            self.assertEqual(client.wait_result(session, request, "retrieve"), [session])

    def test_old_worker_heartbeat_not_ready(self):
        self.assertFalse(client.SERVICE.heartbeat_ready({"timestamp": time.time(), "ready_gpus": 1}))
        heartbeat = {"timestamp": time.time(), "ready_gpus": 1, "visible_gpus": 1,
                     "protocol_version": 2, "capabilities": ["retrieve", "rerank"],
                     "status": "ready", "service_name": "appeals"}
        self.assertTrue(client.SERVICE.heartbeat_ready(heartbeat))
        self.assertEqual(client.SERVICE.num_gpus, 1)


class WorkerInitializationTests(unittest.TestCase):
    def test_main_exits_and_marks_stopped_after_idle_ttl(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            profile = replace(worker.PROFILE, nfs_root=root, idle_timeout_sec=.025)
            with patch.object(worker, "PROFILE", profile), patch.object(
                worker, "activity", worker.WorkerActivity(profile),
            ), patch.object(worker, "NFS_ROOT", root), patch.object(
                worker, "HEARTBEAT", profile.heartbeat_path,
            ), patch.object(worker, "SCAN_INTERVAL_SEC", .005), patch.object(
                worker, "load_runtime", return_value=(Mock(), Mock()),
            ) as load:
                worker.main()
            load.assert_called_once_with()
            heartbeat = json.loads(profile.heartbeat_path.read_text(encoding="utf-8"))
            self.assertEqual(heartbeat["status"], "stopped")
            self.assertEqual(heartbeat["ready_gpus"], 0)

    def test_closed_contour_packages_use_existing_token_without_argv_leak(self):
        with patch.object(worker.osiris_config, "get_package_token", return_value="test@token"), patch.object(
            worker.subprocess, "run", return_value=types.SimpleNamespace(returncode=0),
        ) as install, patch.object(worker.importlib, "invalidate_caches") as invalidate:
            worker.install_runtime_packages()
        args = install.call_args.args[0]
        self.assertIn("sentence-transformers==3.2.0", args)
        self.assertIn("faiss-gpu-cu12", args)
        self.assertNotIn("test@token", repr(args))
        self.assertIn("test%40token", install.call_args.kwargs["env"]["PIP_INDEX_URL"])
        self.assertEqual(install.call_args.kwargs["env"]["PIP_TRUSTED_HOST"], "sberosc.ca.sbrf.ru")
        self.assertTrue(install.call_args.kwargs["capture_output"])
        self.assertNotIn("shell", install.call_args.kwargs)
        invalidate.assert_called_once_with()

    def test_package_token_environment_takes_precedence(self):
        with patch.dict(worker.os.environ, {"TOKEN_OSC": " env-token "}), patch.object(Path, "read_text") as read:
            self.assertEqual(worker.osiris_config.get_package_token(), "env-token")
        read.assert_not_called()

    def test_package_token_reads_skill_env_inside_worker(self):
        with patch.dict(worker.os.environ, {}, clear=True), patch.object(
            Path, "read_text", return_value='# comment\nOTHER=x\nTOKEN_OSC="file-token"\n',
        ) as read:
            self.assertEqual(worker.osiris_config.get_package_token(), "file-token")
        read.assert_called_once_with(encoding="utf-8-sig")

    def test_missing_or_empty_package_token_fails_before_pip(self):
        for content in ("", "TOKEN_OSC=  ", "TOKEN_OSC=''", "# TOKEN_OSC=ignored"):
            with self.subTest(content=content), patch.dict(worker.os.environ, {}, clear=True), patch.object(
                Path, "read_text", return_value=content,
            ), patch.object(worker.subprocess, "run") as install:
                with self.assertRaisesRegex(RuntimeError, "TOKEN_OSC is missing"):
                    worker.install_runtime_packages()
                install.assert_not_called()
        with patch.dict(worker.os.environ, {}, clear=True), patch.object(Path, "read_text", side_effect=FileNotFoundError):
            with self.assertRaisesRegex(RuntimeError, "shared NFS"):
                worker.osiris_config.get_package_token()

    def test_unreadable_secret_has_safe_diagnostic(self):
        with patch.dict(worker.os.environ, {}, clear=True), patch.object(
            Path, "read_text", side_effect=PermissionError("sensitive detail"),
        ):
            with self.assertRaisesRegex(RuntimeError, "Cannot read Appeals Osiris secret file") as caught:
                worker.osiris_config.get_package_token()
            self.assertNotIn("sensitive detail", str(caught.exception))

    def test_pip_failure_does_not_print_credentials_or_continue(self):
        with patch.object(worker.osiris_config, "get_package_token", return_value="secret@value"), patch.object(
            worker.subprocess, "run", return_value=types.SimpleNamespace(
                returncode=1, stdout="secret@value", stderr="secret%40value",
            ),
        ), patch.object(worker.importlib, "invalidate_caches") as invalidate, patch("builtins.print") as output:
            with self.assertRaisesRegex(RuntimeError, "pip exit 1") as caught:
                worker.install_runtime_packages()
        self.assertNotIn("secret", str(caught.exception))
        self.assertNotIn("secret", repr(output.call_args_list))
        invalidate.assert_not_called()

    def test_metadata_loader_does_not_require_texts_or_embeddings(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "bm25s_shards/shard_0").mkdir(parents=True)
            with (root / "meta.pkl").open("wb") as handle:
                pickle.dump({"doc_ids": ["first", "second"], "id_to_index": {"first": 0, "second": 1}}, handle)
            faiss = types.ModuleType("faiss")
            faiss.read_index = Mock(return_value=types.SimpleNamespace(ntotal=2))
            bm25 = types.ModuleType("bm25s")
            bm25.BM25 = types.SimpleNamespace(load=Mock(return_value=types.SimpleNamespace(scores={"num_docs": 2})))
            with patch.dict(sys.modules, {"faiss": faiss, "bm25s": bm25}), patch.object(
                bge, "_loaded", False,
            ), patch.object(bge, "doc_ids", []), patch.object(bge, "req_reg_dates", []), patch.object(
                bge, "id_to_positions", {},
            ), patch.object(bge, "faiss_loaded", None), patch.object(bge, "bm25_indexes", []):
                bge.load_pipeline_meta_and_indices(root)
                self.assertEqual(bge.build_allowed_mask(["first", "missing"]).tolist(), [True, False])
                bm25.BM25.load.assert_called_once_with(str(root / "bm25s_shards/shard_0"), load_corpus=False)

    def test_worker_requires_cuda_without_model_fallback(self):
        torch = types.ModuleType("torch")
        torch.cuda = types.SimpleNamespace(is_available=lambda: False)
        transformers = types.ModuleType("sentence_transformers")
        transformers.CrossEncoder = Mock()
        transformers.SentenceTransformer = Mock()
        with patch.object(worker, "install_runtime_packages"), patch.dict(
            sys.modules, {"torch": torch, "sentence_transformers": transformers},
        ):
            with self.assertRaisesRegex(RuntimeError, "requires CUDA"):
                worker.load_runtime()
        transformers.CrossEncoder.assert_not_called()
        transformers.SentenceTransformer.assert_not_called()

    def test_worker_oom_reduces_batch_and_returns_raw_logits(self):
        torch = types.ModuleType("torch")
        torch.cuda = types.SimpleNamespace(OutOfMemoryError=MemoryError, empty_cache=Mock())
        torch.nn = types.SimpleNamespace(Identity=lambda: "identity")
        model = types.SimpleNamespace(predict=Mock(side_effect=[MemoryError("oom"), np.array([-1., 0., 1.])]))
        with patch.dict(sys.modules, {"torch": torch}):
            raw = worker.predict_with_retry(model, [("query", "text")] * 3)
        self.assertEqual([call.kwargs["batch_size"] for call in model.predict.call_args_list], [4, 2])
        self.assertEqual(model.predict.call_args.kwargs["activation_fct"], "identity")
        self.assertEqual(raw.tolist(), [-1., 0., 1.])
        self.assertEqual(worker.sigmoid(raw)[1], .5)


class NativeDeliveryTests(unittest.IsolatedAsyncioTestCase):
    """Run the real adapter with a minimal offline Nanobot API double."""

    def setUp(self):
        base = types.ModuleType("nanobot.agent.tools.base")
        base.Tool = type("Tool", (), {})
        base.tool_parameters = lambda schema: lambda cls: cls
        base.ToolResult = types.SimpleNamespace(
            error=lambda message: types.SimpleNamespace(is_error=True, message=message),
        )
        context = types.ModuleType("nanobot.agent.tools.context")
        context.current_request_context = lambda: object()
        context.current_request_session_key = lambda: "real/session"
        with patch.dict(sys.modules, {base.__name__: base, context.__name__: context}):
            self.native = load_file("appeals_native_offline", ROOT / "workspace/tools/appeals_analyzer.py")
        self.message = types.SimpleNamespace(execute=AsyncMock(return_value=None))
        self.tool = self.native.AppealsAnalyzerTool(
            config=self.native.AppealsAnalyzerToolConfig(), message_tool=self.message,
        )
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.addCleanup(patch.stopall)
        patch.object(self.native, "log_skill_runtime", return_value="production").start()
        real_import = importlib.import_module
        patch.object(self.native, "import_module", side_effect=lambda name: {
            "appeals_analyzer_runtime.utils.appeals_artifacts": artifacts,
            "appeals_analyzer_runtime.utils.data_store": store,
        }.get(name) or real_import(name)).start()
        patch.object(artifacts, "session_results", side_effect=lambda sid: self.root / sid.replace("/", "_")).start()

    async def test_real_session_cache_backend_and_full_report_media(self):
        async def runner(session_id, user_prompt):
            self.assertEqual(session_id, "real/session")
            self.assertEqual(store._backend.get(), "cache")
            path = artifacts.output_directory(session_id) / "accepted.xlsx"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(b"test artifact")
            artifacts.register_artifact(path)
            return "Полный отчёт\n1\n2\n3\n4"

        with patch.object(self.tool, "_load_runner", return_value=runner), store.backend_scope("greenplum"):
            result = await self.tool.execute(prompt=PROMPT, session_id="forged")
            self.assertEqual(store._backend.get(), "greenplum")
        self.message.execute.assert_awaited_once_with(
            content=result, media=[str(self.root / "real_session/accepted.xlsx")],
        )

    async def test_request_without_session_fails_before_pipeline(self):
        runner = AsyncMock()
        with patch.object(self.tool, "_load_runner", return_value=runner), patch.object(
            self.native, "current_request_session_key", return_value=None,
        ):
            result = await self.tool.execute(prompt=PROMPT)
        self.assertTrue(result.is_error)
        runner.assert_not_awaited()
        self.message.execute.assert_not_awaited()

    async def test_real_loader_needs_neither_gp_nor_ml(self):
        foreign = types.ModuleType("utils")
        foreign.__path__ = ["foreign-utils"]
        with patch.dict(sys.modules, {"utils": foreign}):
            runner = self.tool._load_runner()
        self.assertEqual(runner.__module__, "appeals_analyzer_runtime.scripts.appeals_reports")
        self.assertNotIn("appeals_analyzer_runtime.utils.db", sys.modules)
        self.assertNotIn("appeals_analyzer_runtime.utils.bge_search_engine", sys.modules)

    async def test_worker_startup_error_reaches_user_as_retry_message(self):
        self.tool._load_runner()
        config = importlib.import_module("appeals_analyzer_runtime.utils.osiris_config")
        runner = AsyncMock(side_effect=config.OsirisWorkerNotReady(
            "Appeals Osiris worker is not ready. Start it with: python appeals_osiris_job.py start",
        ))
        with patch.object(self.tool, "_load_runner", return_value=runner):
            result = await self.tool.execute(prompt=PROMPT)
        self.assertTrue(result.is_error)
        self.assertIn("через 15 минут", result.message)
        self.message.execute.assert_not_awaited()

    async def test_empty_result_has_no_media(self):
        runner = AsyncMock(return_value="Релевантные обращения не подтверждены.")
        with patch.object(self.tool, "_load_runner", return_value=runner):
            result = await self.tool.execute(prompt=PROMPT)
        self.assertIn("не подтверждены", result)
        self.message.execute.assert_not_awaited()

    async def test_artifact_scopes_are_concurrent_and_do_not_scan_directories(self):
        async def run(sid):
            paths = []
            target = self.root / sid
            target.mkdir()
            (target / "old.xlsx").write_bytes(b"old")
            with artifacts.artifact_scope(target, paths):
                await asyncio.sleep(0)
                path = artifacts.output_directory(sid) / "new.xlsx"
                path.write_bytes(b"new")
                artifacts.register_artifact(path)
            return paths
        left, right = await asyncio.gather(run("left"), run("right"))
        self.assertEqual(left, [self.root / "left/new.xlsx"])
        self.assertEqual(right, [self.root / "right/new.xlsx"])

    async def test_testing_branch_does_not_load_production(self):
        runner = types.SimpleNamespace(run_testing_report=AsyncMock(return_value="testing report"))
        with patch.object(self.native, "log_skill_runtime", return_value="testing"), patch.object(
            self.native, "current_tool_session_id", return_value="testing-session",
        ), patch.object(self.native, "load_testing_module", return_value=runner), patch.object(
            self.tool, "_load_runner", side_effect=AssertionError("production loaded"),
        ):
            self.assertEqual(await self.tool.execute(prompt="test"), "testing report")
        runner.run_testing_report.assert_awaited_once_with(session_id="testing-session", user_prompt="test")


class StandaloneBackendTests(unittest.TestCase):
    def test_cli_owns_gp_scope_and_stops_pool_on_success_and_failure(self):
        cli = load_file("appeals_cli_offline", SKILL / "scripts/cli.py")
        for fails in (False, True):
            with self.subTest(fails=fails):
                args = types.SimpleNamespace(
                    prompt='"", "", "", "query"', prompt_file=None, positional_prompt=[],
                    profile="test", session_id="cli-session", allow_vllm=False,
                    log_level="INFO", console_only=True, log_file=None,
                )
                db = types.SimpleNamespace(shutdown=Mock())
                config = types.ModuleType("config")
                config.is_settings_initialized = lambda: False
                config._initialize_settings = Mock()
                async def runner(**kwargs):
                    self.assertEqual(store._backend.get(), "greenplum")
                    if fails:
                        raise RuntimeError("expected CLI failure")
                    return "report"
                with patch.dict(sys.modules, {"config": config}), patch.object(
                    cli, "build_parser", return_value=types.SimpleNamespace(parse_args=lambda: args),
                ), patch.object(cli, "configure_logging"), patch.object(cli, "load_shared_db", return_value=db), patch.object(
                    cli, "start_standalone_db_runtime",
                ) as start, patch.object(cli, "load_standalone_runner", return_value=runner), patch.object(
                    cli.importlib, "import_module", return_value=store,
                ), patch.dict(cli.os.environ, {}, clear=False):
                    self.assertEqual(cli.main(), int(fails))
                start.assert_called_once_with(db)
                db.shutdown.assert_called_once_with()
                config._initialize_settings.assert_called_once_with("test")
                self.assertEqual(store._backend.get(), "cache")


if __name__ == "__main__":
    unittest.main()

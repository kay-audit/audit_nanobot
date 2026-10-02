"""Извлечение текста из документов без изменения оригинала."""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable


SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".pptx", ".docx"}
PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "raw" / "extracted"


class ExtractionError(RuntimeError):
    """Понятная пользователю ошибка извлечения."""


@dataclass
class Extraction:
    text: str
    warnings: list[str]


def extract_text_file(source: Path) -> Extraction:
    errors: list[str] = []
    for encoding in ("utf-8-sig", "utf-16", "cp1251"):
        try:
            text = source.read_text(encoding=encoding)
            warnings = []
            if encoding != "utf-8-sig":
                warnings.append(
                    f"Файл прочитан в кодировке {encoding}; проверьте символы."
                )
            return Extraction(text, warnings)
        except UnicodeError as exc:
            errors.append(f"{encoding}: {exc}")
    raise ExtractionError(
        "Не удалось определить кодировку. Попытки: " + "; ".join(errors)
    )


def extract_pdf(source: Path) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError(
            "Установите pypdf: python3.12 -m pip install pypdf"
        ) from exc

    try:
        reader = PdfReader(str(source))
    except Exception as exc:
        raise ExtractionError(f"Не удалось открыть PDF: {exc}") from exc

    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise ExtractionError("PDF зашифрован и требует пароль.") from exc
        if unlocked == 0:
            raise ExtractionError("PDF зашифрован и требует пароль.")

    sections: list[str] = []
    empty_pages: list[int] = []
    for number, page in enumerate(reader.pages, 1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            text = f"> Ошибка извлечения страницы: {exc}"
            empty_pages.append(number)
        if not text:
            text = "> Текст не извлечён; возможно, страница является сканом."
            empty_pages.append(number)
        sections.append(f"## Страница {number}\n\n{text}")

    warnings = [
        "Извлечён только доступный текст; верстка, изображения, формулы и часть "
        "таблиц могли быть потеряны.",
        "OCR не выполнялся; сканированные страницы могут не содержать текста.",
    ]
    if empty_pages:
        warnings.append(
            "Страницы без надёжно извлечённого текста: "
            + ", ".join(map(str, sorted(set(empty_pages))))
            + "."
        )
    return Extraction("\n\n".join(sections), warnings)


def extract_pptx(source: Path) -> Extraction:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ExtractionError(
            "Установите python-pptx: python3.12 -m pip install python-pptx"
        ) from exc

    try:
        presentation = Presentation(str(source))
    except Exception as exc:
        raise ExtractionError(f"Не удалось открыть PPTX: {exc}") from exc

    slides: list[str] = []
    empty_slides: list[int] = []
    for number, slide in enumerate(presentation.slides, 1):
        parts = [f"## Слайд {number}"]
        title = slide.shapes.title
        title_id = title.shape_id if title is not None else None
        found = False

        if title is not None and title.has_text_frame and title.text.strip():
            parts.append(f"### {title.text.strip()}")
            found = True

        for shape in slide.shapes:
            if title_id is not None and shape.shape_id == title_id:
                continue
            if getattr(shape, "has_text_frame", False):
                paragraphs = [
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                ]
                if paragraphs:
                    parts.append("\n".join(f"- {item}" for item in paragraphs))
                    found = True
            if getattr(shape, "has_table", False):
                rows = [
                    " | ".join(
                        cell.text.strip().replace("\n", " ") for cell in row.cells
                    )
                    for row in shape.table.rows
                ]
                if rows:
                    parts.append("```text\n" + "\n".join(rows) + "\n```")
                    found = True

        if not found:
            parts.append("> На слайде не найден извлекаемый текст.")
            empty_slides.append(number)
        slides.append("\n\n".join(parts))

    warnings = [
        "Извлечены текстовые блоки и таблицы; оформление, изображения, диаграммы "
        "и расположение могли быть потеряны.",
        "Заметки докладчика и встроенные объекты не извлекаются.",
    ]
    if empty_slides:
        warnings.append(
            "Слайды без извлекаемого текста: "
            + ", ".join(map(str, empty_slides))
            + "."
        )
    return Extraction("\n\n".join(slides), warnings)


def extract_docx(source: Path) -> Extraction:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ExtractionError(
            "Установите python-docx: python3.12 -m pip install python-docx"
        ) from exc

    try:
        document = Document(str(source))
    except Exception as exc:
        raise ExtractionError(f"Не удалось открыть DOCX: {exc}") from exc

    parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            if style.startswith("Heading"):
                try:
                    level = int(style.split()[-1])
                except (ValueError, IndexError):
                    level = 1
                parts.append(f"{'#' * min(max(level + 1, 2), 6)} {text}")
            elif style.lower().startswith(("list", "список")):
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows = [
                " | ".join(
                    cell.text.strip().replace("\n", " ") for cell in row.cells
                )
                for row in table.rows
            ]
            if rows:
                parts.append("```text\n" + "\n".join(rows) + "\n```")

    warnings = [
        "Извлечены основной текст и таблицы; форматирование, колонтитулы, "
        "комментарии, изображения, сноски и поля могли быть потеряны."
    ]
    if not parts:
        warnings.append("В документе не найден извлекаемый основной текст.")
    return Extraction("\n\n".join(parts), warnings)


def extract_document(source: Path) -> Extraction:
    """Извлечь поддерживаемый документ в памяти, ничего не записывая."""

    suffix = source.suffix.lower()
    extractors: dict[str, Callable[[Path], Extraction]] = {
        ".txt": extract_text_file,
        ".md": extract_text_file,
        ".markdown": extract_text_file,
        ".pdf": extract_pdf,
        ".pptx": extract_pptx,
        ".docx": extract_docx,
    }
    extractor = extractors.get(suffix)
    if extractor is None:
        raise ExtractionError(
            "Формат не поддерживается. Допустимы: "
            + ", ".join(sorted(SUPPORTED))
        )
    return extractor(source)


def source_display_path(
    source: Path,
    *,
    project_root: Path = PROJECT_ROOT,
) -> str:
    resolved = source.resolve()
    try:
        return resolved.relative_to(project_root.resolve()).as_posix()
    except ValueError:
        return str(resolved)


def yaml_string(value: str) -> str:
    # JSON-строка допустима в YAML и не требует зависимости PyYAML.
    return json.dumps(value, ensure_ascii=False)


def render(
    source: Path,
    extraction: Extraction,
    *,
    original_path: str | None = None,
    original_sha256: str | None = None,
) -> str:
    extracted_at = datetime.now().astimezone().isoformat(timespec="seconds")
    lines = [
        "---",
        "original_path: "
        + yaml_string(original_path or source_display_path(source)),
        f"file_type: {yaml_string(source.suffix.lstrip('.').upper())}",
        f"extracted_at: {yaml_string(extracted_at)}",
    ]
    if original_sha256:
        lines.append(
            f"original_sha256: {yaml_string(original_sha256)}"
        )
    if extraction.warnings:
        lines.append("warnings:")
        lines.extend(f"  - {yaml_string(item)}" for item in extraction.warnings)
    else:
        lines.append("warnings: []")
    lines.extend(
        [
            "---",
            "",
            "# Извлечённый текст",
            "",
            "> Это вспомогательная копия. Первоисточником остаётся файл из "
            "`original_path`.",
            "",
            extraction.text.rstrip(),
            "",
        ]
    )
    return "\n".join(lines)


def write_output(target: Path, content: str, force: bool) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() and not force:
        raise ExtractionError(
            f"Результат уже существует: {target}. "
            "Используйте --force только для обновления извлечённой копии."
        )

    temporary = target.with_name(f".{target.name}.tmp")
    if temporary.exists():
        raise ExtractionError(f"Временный файл уже существует: {temporary}")
    try:
        temporary.write_text(content, encoding="utf-8")
        temporary.replace(target)
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        raise ExtractionError(f"Не удалось записать результат: {exc}") from exc


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Извлечь TXT, Markdown, PDF, PPTX или DOCX в raw/extracted/, "
            "не изменяя исходный файл."
        )
    )
    parser.add_argument("source", type=Path, help="Исходный документ")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Каталог результата (по умолчанию raw/extracted/)",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Перезаписать только существующую извлечённую копию",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source = args.source.expanduser()
    if not source.exists() or not source.is_file():
        print(f"Ошибка: исходный файл не найден: {source}", file=sys.stderr)
        return 2

    suffix = source.suffix.lower()
    if suffix not in SUPPORTED:
        print(
            "Ошибка: формат не поддерживается. Допустимы: "
            + ", ".join(sorted(SUPPORTED)),
            file=sys.stderr,
        )
        return 2

    target = args.output_dir.expanduser() / f"{source.name}.md"

    try:
        if source.resolve() == target.resolve():
            raise ExtractionError("Путь результата совпадает с оригиналом.")
        extraction = extract_document(source)
        write_output(target, render(source, extraction), args.force)
    except (ExtractionError, OSError) as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 1

    print(f"Готово: {target}")
    if extraction.warnings:
        print("Предупреждения:")
        for warning in extraction.warnings:
            print(f"- {warning}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

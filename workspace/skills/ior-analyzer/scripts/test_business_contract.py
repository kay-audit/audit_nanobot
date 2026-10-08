"""End-to-end deterministic fixtures: business outcomes, not SQL substrings."""
import asyncio
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import AsyncMock, patch
import unittest
import duckdb
import pandas as pd

import ior_reports as reports
from utils.data_store import DUCKDB_TABLES
from utils.resolve.money_parser import parse_money
from utils.resolve.period_parser import parse_period
from utils.resolve.request_plan import build_plan, resolve_preset, explicit_followup
from utils.resolve.request_outcome import ClarificationRequired
from utils.resolve import semantic_grounding as ground
from utils import bge_search_engine as bge
from utils.session_extract_manager import set_session_extract,get_session_extract
from preset_analysis.slices import render_slices
import pytest

CAT={ground.ORIGIN:['Среднерусский банк','Московская региональная дирекция','Риски','Люди и культура'],
     ground.RESPONSIBILITY:['Риски','Люди и культура']}
def frame():
    df = pd.DataFrame({'incdnt_id':[1,2,3],'incdnt_sid':['EVE-1','EVE-2','EVE-3'],
        'incdnt_entry_dt':pd.to_datetime(['2025-01-01','2025-05-01','2025-08-01']),
        'incdnt_status_name':['Утверждён']*3,'org_struct_id':['SBR_A']*3,
        ground.ORIGIN:['Среднерусский банк','Московская региональная дирекция','Риски'],
        'org_struct_lvl_4_name':['Другой','Среднерусский банк','Среднерусский банк'],
        ground.RESPONSIBILITY:['Риски','Риски','Люди и культура'],
        'risk_profile_id':['DRP-1','DRP-2','DRP-1'],'risk_profile_name':['А','Б','А'],
        'incdnt_sum':[Decimal('1500000'),Decimal('1000000'),Decimal('2000000')],
        'recovery_rub_amt_aggr':[0.,0.,0.],
        'incdnt_full_descr_txt':['Первый','Второй','Третий']})
    for c in reports.MAIN_REPORT_COLUMNS:
        if c not in df: df[c] = None
    return df

@pytest.mark.parametrize('text,expected', [
    ('первый и третий кварталы 2025', [('2025-01-01','2025-04-01'),('2025-07-01','2025-10-01')]),
    ('за первый, второй и четвертый кварталы 2025', [('2025-01-01','2025-04-01'),('2025-04-01','2025-07-01'),('2025-10-01','2026-01-01')]),
    ('с первого по третий квартал 2025', [('2025-01-01','2025-10-01')]),
    ('последние три месяца', [('2026-07-08','2026-10-09')]),
    ('последние пять дней', [('2026-10-04','2026-10-09')]),
    ('последние двенадцать месяцев', [('2025-10-08','2026-10-09')]),
    ('с 15 марта 2025 по 20 апреля 2025', [('2025-03-15','2025-04-21')]),
    ('с 15-го марта 2025 по 20-е апреля 2025 года', [('2025-03-15','2025-04-21')]),
    ('15 марта 2025', [('2025-03-15','2025-03-16')]),
    ('с 31 декабря 2024 по 2 января 2025', [('2024-12-31','2025-01-03')]),
])
def test_second_pass_periods(text,expected):
    assert parse_period(text,today=date(2026,10,8)).intervals == expected

@pytest.mark.parametrize('text', ['с 31 февраля 2025 по 20 марта 2025', 'с 20 апреля 2025 по 15 марта 2025'])
def test_second_pass_invalid_named_dates(text):
    with pytest.raises(ClarificationRequired): parse_period(text)

@pytest.mark.parametrize('text,expected', [
    ('больше полутора миллионов', [('>', '1500000.0')]),
    ('больше одного миллиона', [('>', '1000000')]),
    ('от миллиона до двух миллионов', [('>=','1000000'),('<=','2000000')]),
    ('не менее одной тысячи', [('>=','1000')]),
    ('не более трёх миллионов', [('<=','3000000')]),
    ('меньше трех миллионов', [('<','3000000')]),
    ('более четырёх миллиардов', [('>','4000000000')]),
    ('менее четырех тысяч', [('<','4000')]),
    ('от пяти тысяч до десяти миллионов', [('>=','5000'),('<=','10000000')]),
    ('от двух до пяти миллионов', [('>=','2000000'),('<=','5000000')]),
    ('более чем на 1 млрд рублей', [('>','1000000000')]),
    ('больше чем на 1,5 млн рублей', [('>','1500000.0')]),
])
def test_second_pass_money(text,expected):
    assert [(m.op,m.value) for m in parse_money(text)] == [(op,Decimal(n)) for op,n in expected]

@pytest.mark.parametrize('text,preset', [
    ('ИОРы с суммой больше миллиона','financial_consequences_ior'),
    ('ИОРы более чем на 1 млрд рублей','financial_consequences_ior'),
    ('ИОРы больше чем на 1 млрд рублей','financial_consequences_ior'),
    ('потери больше миллиона','financial_consequences_ior'),
    ('ИОРы с потерями от 500 тысяч до 2 миллионов','financial_consequences_ior'),
    ('возмещения больше миллиона','vozmeshenie_ior'),
    ('сумма возмещения более миллиона','vozmeshenie_ior'),
    ('ИОРы с суммой возмещения больше 1 млн','vozmeshenie_ior'),
    ('возмещено более 500 тысяч','vozmeshenie_ior'),
    ('возмещения потерь больше миллиона','vozmeshenie_ior'),
])
def test_second_pass_money_routing(text,preset):
    with patch.object(ground,'load_catalog',return_value=CAT):
        plan=build_plan('ior_hypothesis',text)
    assert plan.preset == preset
    assert plan.money
    assert not any('incdnt_sum' in p for p in plan.predicates())

@pytest.mark.parametrize('reference', ['в этой выборке','из этой выборки','в полученной выборке',
    'из полученной выборки','в предыдущем отчёте','в прошлой выгрузке','в последней выгрузке',
    'среди выгруженных выше','среди этих ИОР','из этих ИОР'])
def test_second_pass_explicit_followup(reference):
    assert explicit_followup('Расскажи подробнее про EVE-1 '+reference)
    assert not explicit_followup('Найди EVE-1')

@pytest.mark.parametrize('text', ['прямые потери больше миллиона','косвенные потери более 1 млн'])
def test_second_pass_concrete_losses_still_require_tab(text):
    with pytest.raises(ClarificationRequired,match='вкладк'): resolve_preset(None,text)

class BusinessTests(unittest.TestCase):
    def setUp(self):
        self.catalog=patch.object(ground,'load_catalog',return_value=CAT)
        self.catalog.start()
        self.df=frame()
        self.con=duckdb.connect(':memory:')
        self.con.register(DUCKDB_TABLES['ior'],self.df)
        # Monetary chat queries now use consequences, independently of main-table sums.
        import re
        sql=reports.build_preset_sql_queries(DUCKDB_TABLES)['financial_consequences_ior']
        cols=set(re.findall(r'\bfi\.(\w+)',sql))
        details=pd.DataFrame({c:[None]*3 for c in cols})
        details['incdnt_id']=[1,2,3]
        details['fin_impact_sid']=['FI-1','FI-2','FI-3']
        details['fin_impact_rub_amt']=[1500000.,1000000.,2000000.]
        self.con.register(DUCKDB_TABLES['financial_impact'],details)
    def tearDown(self):
        self.con.close()
        self.catalog.stop()
    def select(self,prompt,preset=None):
        return self.con.execute(reports.build_dynamic_sql_from_prompt(prompt,preset_name=preset,tables=DUCKDB_TABLES)).df()
    def test_origin_only_lvl3(self):
        self.assertEqual(self.select('Выведи ИОРы по СРБ за 2025 год').incdnt_sid.tolist(),['EVE-1'])
    def test_ambiguity_and_business_answers(self):
        for prompt in ['ИОРы по блоку Риски','ИОРы по СРБ и блоку Риски']:
            with self.assertRaises(ClarificationRequired): self.select(prompt)
        origin=self.select('ИОРы по СРБ и блоку Риски за 2025. Уточнение: Риски — место происхождения')
        self.assertEqual(set(origin.incdnt_sid),{'EVE-1','EVE-3'})
        responsible=self.select('ИОРы по СРБ и блоку Риски за 2025. Уточнение: Риски — зона ответственности')
        self.assertEqual(set(responsible.incdnt_sid),{'EVE-1'})
    def test_single_category_catalog(self):
        for column,category in [(ground.ORIGIN,'origin'),(ground.RESPONSIBILITY,'responsibility')]:
            with patch.object(ground,'load_catalog',return_value={column:['Риски']}):
                plan=build_plan(None,'ИОРы по блоку Риски')
                self.assertEqual([h['category'] for h in plan.hits],[category])
    def test_responsibility_values_are_or(self):
        result=self.select('ИОРы по блоку Риски и Люди и культура. Уточнение: зона ответственности')
        self.assertEqual(set(result.incdnt_sid),{'EVE-1','EVE-2','EVE-3'})

    def test_separate_answers_for_two_ambiguous_blocks(self):
        plan=build_plan(None,'ИОРы по Риски и Люди и культура. Уточнение: Риски — место происхождения; Люди и культура — зона ответственности')
        self.assertEqual({(h['category'],h['value']) for h in plan.hits},{('origin','Риски'),('responsibility','Люди и культура')})
        with self.assertRaises(ClarificationRequired):
            build_plan(None,'ИОРы по Риски и Люди и культура. Уточнение: Риски — место происхождения')
    def test_codes_and_routing(self):
        self.assertEqual(resolve_preset(None,'финансовые последствия EVE-1'),'financial_consequences_ior')
        self.assertEqual(resolve_preset(None,'нефинансовые последствия EVE-1'),'ior_nonfinancial_consequences')
        self.assertEqual(resolve_preset(None,'финансовые последствия ПАО Сбербанк'),'financial_consequences_ior')
        self.assertEqual(resolve_preset(None,'Досье EVE-1'),'report_period_specific_ior')
        with self.assertRaises(ClarificationRequired): build_plan(None,'Досье EVE-1 и EVE-2')
        self.assertEqual(set(self.select('ИОРы DRP-1 и DRP-2').incdnt_sid),{'EVE-1','EVE-2','EVE-3'})
        for prompt in ['ИОРы SBR-123','Возмещения по удалённым ИОР','Прямые потери больше миллиона']:
            with self.assertRaises(ClarificationRequired): self.select(prompt)
    def test_decimal_boundaries(self):
        for wording in ['более 1,5 млн','более 1.5 млн']:
            self.assertEqual(self.select('ИОРы с суммой '+wording).incdnt_sid.tolist(),['EVE-3'])
        self.assertEqual(set(self.select('ИОРы с суммой не менее 1,5 млн').incdnt_sid),{'EVE-1','EVE-3'})
        self.assertEqual(self.select('ИОРы с суммой менее 1 млн').incdnt_sid.tolist(),[])
        self.assertEqual(self.select('ИОРы с суммой не более 1 млн').incdnt_sid.tolist(),['EVE-2'])
        self.assertEqual(len(self.select('ИОРы с суммой от 1 млн до 2 млн')),3)
    def test_monetary_default_does_not_use_main_sum(self):
        self.df['incdnt_sum']=[0.,9000000000.,0.]
        for text in ['ИОРы с суммой больше миллиона', 'потери больше миллиона',
                     'ИОРы с потерями от 1,5 млн до 2 миллионов']:
            with self.subTest(text=text):
                plan=build_plan(None,text)
                self.assertEqual(plan.preset,'financial_consequences_ior')
                self.assertEqual(set(self.select(text).incdnt_sid),{'EVE-1','EVE-3'})
                self.assertFalse(any('incdnt_sum' in c for c in plan.predicates()))
    def test_amount_forms(self):
        for text,value in [('500 тысяч',500000),('1 млрд',1000000000),('2 миллиарда',2000000000),('полтора миллиона',1500000),('милион',1000000),('1500 тыс',1500000),('1 500 000',1500000)]:
            with self.subTest(text=text):
                conditions=parse_money('больше '+text)
                self.assertEqual((conditions[0].op,conditions[0].value),('>',Decimal(value)))
        with self.assertRaises(ClarificationRequired): self.select('ИОРы с суммой больше нескольких миллионов')
    def test_periods_and_slices(self):
        selected=self.select('ИОРы за Q1 и Q3 2025')
        self.assertEqual(set(selected.incdnt_sid),{'EVE-1','EVE-3'})
        plan=build_plan(None,'ИОРы за Q1 и Q3 2025 по DRP-1 и DRP-2')
        text=render_slices(selected,plan)
        for label in ['Q1 2025','Q3 2025','DRP-1','DRP-2']: self.assertIn(label,text)
        self.assertEqual(len(self.select('ИОРы за Q1-Q3 2025')),3)
        cases=[('за 2024 и 2025 годы',[('2024-01-01','2025-01-01'),('2025-01-01','2026-01-01')]),
               ('декабрь 2024 — январь 2025',[('2024-12-01','2025-02-01')]),
               ('с 15.03.2025 по 20.04.2025',[('2025-03-15','2025-04-21')]),
               ('прошлый год',[('2025-01-01','2026-01-01')]),
               ('прошлый месяц',[('2026-09-01','2026-10-01')]),
               ('последние 3 месяца',[('2026-07-08','2026-10-09')])]
        for prompt,expected in cases:
            with self.subTest(prompt=prompt): self.assertEqual(parse_period(prompt,today=date(2026,10,8)).intervals,expected)
        with self.assertRaises(ClarificationRequired): self.select('ИОРы за прошлый неизвестный период')
    def test_no_truncation(self):
        self.con.unregister(DUCKDB_TABLES['ior'])
        large=pd.concat([self.df.iloc[:1]]*100001,ignore_index=True)
        large['incdnt_sid']=['EVE-'+str(i) for i in range(len(large))]
        self.con.register(DUCKDB_TABLES['ior'],large)
        self.assertEqual(len(self.select('ИОРы за 2025 год')),100001)
    def test_monetary_subject_sources(self):
        for preset,prompt,logical,prefix,amount in [('financial_consequences_ior','финансовые последствия больше 1 млн','financial_impact','fin_impact','fin_impact_rub_amt'),('vozmeshenie_ior','возмещения больше миллиона','recovery','recovery','recovery_rub_amt')]:
            with self.subTest(preset=preset):
                queries=reports.build_preset_sql_queries(DUCKDB_TABLES)
                # Build the full physical detail projection of the actual preset.
                import re
                alias='fi' if logical=='financial_impact' else 'r'
                columns=set(re.findall(r'\b'+alias+r'\.(\w+)',queries[preset]))
                details=pd.DataFrame({c:[None]*3 for c in columns})
                details['incdnt_id']=[1,1,2]
                details[prefix+'_sid']=['A','B','C']
                details[amount]=[600000.,600000.,1000000.]
                self.con.register(DUCKDB_TABLES[logical],details)
                actual=self.select(prompt,preset)
                self.assertEqual(set(actual.incdnt_sid),{'EVE-1'})
                self.assertEqual(actual[amount].sum(),1200000.)
    def test_catalog_failure_is_closed(self):
        with patch.object(ground,'load_catalog',side_effect=ClarificationRequired('справочник недоступен')):
            with self.assertRaises(ClarificationRequired): self.select('ИОРы по блоку Риски')

    def test_multiple_eve_answer_and_subject_priority(self):
        chosen=build_plan(None,'Досье EVE-1 и EVE-2. Уточнение: EVE-2')
        self.assertEqual([(h['category'],h['value']) for h in chosen.hits],[('eve','EVE-2')])
        for subject,preset in [('возмещения','vozmeshenie_ior'),('финансовые последствия','financial_consequences_ior'),('нефинансовые последствия','ior_nonfinancial_consequences')]:
            self.assertEqual(resolve_preset(None,subject+' EVE-1 и EVE-2'),preset)
        self.assertEqual(resolve_preset(None,'Возмещения по удалённым ИОР. Уточнение: возмещения'),'vozmeshenie_ior')

    def test_single_dossier_executes_main_financial_recovery_join(self):
        import re
        from preset_analysis.report_period_specific_ior_analysis import prepare_dossier_views
        sql=reports.build_preset_sql_queries(DUCKDB_TABLES)['report_period_specific_ior']
        for alias,logical,sid,amount,values in [('fi','financial_impact','fin_impact_sid','fin_impact_rub_amt',[10.,20.]),('r','recovery','recovery_sid','recovery_rub_amt',[1.,2.])]:
            cols=set(re.findall(r'\b'+alias+r'\.(\w+)',sql))
            child=pd.DataFrame({c:[None,None] for c in cols})
            child['incdnt_id']=[1,1]; child[sid]=[alias+'1',alias+'2']; child[amount]=values
            self.con.register(DUCKDB_TABLES[logical],child)
        actual=self.select('Подробные сведения по EVE-1')
        self.assertEqual(set(actual.incdnt_sid),{'EVE-1'})
        self.assertEqual(len(actual),4)
        metrics=prepare_dossier_views(actual)
        self.assertEqual((metrics['fin_count'],metrics['fin_total'],metrics['recovery_count'],metrics['recovery_total']),(2,30.,2,3.))

    def test_explicit_credit_source_and_missing_source_failure(self):
        credits=pd.DataFrame({'incdnt_id':[1],'credit_agr_num':['AGR-1'],'credit_debt_rub_amt':[100.]})
        self.con.register(DUCKDB_TABLES['credits'],credits)
        actual=self.select('Кредитная задолженность по СРБ за 2025','credit_no_way_collect_debt')
        self.assertEqual(actual.credit_agr_num.tolist(),['AGR-1'])
        with self.assertRaisesRegex(RuntimeError,'physical source'):
            reports.build_dynamic_sql_from_prompt('Кредитная задолженность по СРБ за 2025',preset_name='credit_no_way_collect_debt',tables=reports.GREENPLUM_TABLES)

    def test_excel_null_is_blank_and_zero_is_numeric(self):
        import tempfile
        import openpyxl
        from utils.dataframe_ops import prepare_df_for_excel
        from utils.excel_inspector import inspect_excel
        raw=frame().iloc[:2].copy(); raw['incdnt_sum']=[None,0.]
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'money.xlsx'
            prepare_df_for_excel(raw).to_excel(path,index=False,engine='openpyxl')
            book=openpyxl.load_workbook(path); sheet=book.active
            col=[cell.value for cell in sheet[1]].index('incdnt_sum')+1
            self.assertIsNone(sheet.cell(2,col).value)
            self.assertEqual(sheet.cell(3,col).value,0)
            self.assertEqual(sheet.cell(3,col).data_type,'n')
            book.close()
            self.assertEqual(inspect_excel(path)['stats']['sum_total_loss'],0.)

    def test_invalid_period_and_money_are_closed(self):
        for prompt in ['ИОРы за 31.02.2025','ИОРы с 31.02.2025 по 01.03.2025','ИОРы за Q1 без года','ИОРы с суммой более 1,abc млн','ИОРы по SBR','ИОРы с суммой от 2 млн до 1 млн']:
            with self.subTest(prompt=prompt),self.assertRaises(ClarificationRequired): self.select(prompt)

    def test_nullable_amounts_and_zero(self):
        from preset_analysis.common import collapse_detail_to_incidents,format_amount
        from preset_analysis.registry import get_analyzer
        raw=frame().iloc[:2].copy()
        raw['incdnt_sum']=[None,0.]
        bundle=get_analyzer('ior_hypothesis').prepare(raw.iloc[:1])
        self.assertIsNone(bundle.full_metrics['consequences'])
        self.assertEqual(format_amount(None),'')
        self.assertEqual(format_amount(0),'0.00 ₽')
        collapsed=collapse_detail_to_incidents(raw,amount_columns=('incdnt_sum',))
        self.assertTrue(pd.isna(collapsed.iloc[0].incdnt_sum))
        self.assertEqual(collapsed.iloc[1].incdnt_sum,0.)

    def test_deleted_left_join_gap_is_not_action(self):
        from preset_analysis.registry import get_analyzer
        raw=frame().iloc[:2].assign(incdnt_status_name='Удалён',stts_chng_action_dttm=[pd.Timestamp('2025-01-01'),pd.NaT],stts_chng_comment_txt=['Дубликат',None])
        self.con.unregister(DUCKDB_TABLES['ior']); self.con.register(DUCKDB_TABLES['ior'],raw)
        status=pd.DataFrame({'incdnt_id':[1],'incdnt_status_name':['Черновик'],'incdnt_status_code':['draft'],'stts_chng_action_code':['delete'],'stts_chng_action_name':['Удалить'],'stts_chng_comment_txt':['Дубликат'],'stts_chng_action_dttm':pd.to_datetime(['2025-01-01']),'stts_chng_user_num':[None]})
        self.con.register(DUCKDB_TABLES['status'],status)
        actual=self.select('Удалённые ИОР за 2025')
        self.assertEqual(set(actual.incdnt_sid),{'EVE-1','EVE-2'})
        self.assertTrue(pd.isna(actual.set_index('incdnt_sid').loc['EVE-2','stts_chng_action_dttm']))
        self.assertTrue(pd.isna(actual.set_index('incdnt_sid').loc['EVE-2','stts_chng_comment_txt']))
        bundle=get_analyzer('deleted_ior').prepare(actual)
        self.assertEqual(bundle.full_metrics['unique_incidents'],2)
        self.assertEqual(bundle.full_metrics['journal_rows'],1)
        self.assertNotIn('без истории',bundle.deterministic_report())

    def test_actual_plan_context_retains_all_constraints(self):
        from ior_hypothesis import build_analysis_context_text
        plan=build_plan(None,'Финансовые последствия по СРБ и DRP-1 за Q1 и Q3 2025 больше миллиона')
        context=plan.context('исходный запрос')
        self.assertEqual(context['constraints']['origin'],['Среднерусский банк'])
        self.assertEqual(context['constraints']['drp'],['DRP-1'])
        self.assertEqual(len(context['period_intervals']),2)
        self.assertEqual(context['money'],[('>','1000000')])
        rendered=build_analysis_context_text('исходный запрос',context)
        self.assertIn('Среднерусский банк',rendered)
        self.assertIn('не',rendered.lower())

    def test_child_join_population_is_not_limited(self):
        import re
        columns=set(re.findall(r'\bfi\.(\w+)',reports.build_preset_sql_queries(DUCKDB_TABLES)['financial_consequences_ior']))
        details=pd.DataFrame({c:[None]*100001 for c in columns})
        details['incdnt_id']=1; details['fin_impact_sid']=['F'+str(i) for i in range(len(details))]; details['fin_impact_rub_amt']=1.
        self.con.register(DUCKDB_TABLES['financial_impact'],details)
        actual=self.select('Финансовые последствия по СРБ за 2025')
        self.assertEqual(len(actual),100001)
        self.assertEqual(actual.fin_impact_rub_amt.sum(),100001.)

    def test_explicit_projection_and_no_population_limit(self):
        for sql in reports.build_preset_sql_queries(DUCKDB_TABLES).values():
            self.assertNotIn('SELECT *',sql)
            self.assertNotIn('ior.*',sql)
            self.assertNotIn('LIMIT',sql.upper())

    def test_internal_query_adapter_uses_registry_and_halfopen_or(self):
        from types import SimpleNamespace
        from utils.query_adapter import select_sql
        definition=SimpleNamespace(column_names=lambda:list(self.df.columns))
        schema=SimpleNamespace(get=lambda name:definition)
        store=SimpleNamespace(tables=DUCKDB_TABLES)
        sql=select_sql(store,schema,DUCKDB_TABLES['ior'],parse_period('Q1 и Q3 2025').as_filter(),['incdnt_sid'],'incdnt_sid')
        self.assertEqual(set(self.con.execute(sql).df().incdnt_sid),{'EVE-1','EVE-3'})
        self.assertNotIn('LIMIT',sql)
        with self.assertRaises(ValueError): select_sql(store,schema,DUCKDB_TABLES['ior'],{'evil; DROP TABLE x':1},['incdnt_sid'])

    def test_date_ranges_do_not_become_money_and_status_is_retained(self):
        self.assertFalse(build_plan(None,'ИОРы от 2024 до 2025 года').money)
        self.assertEqual(parse_period('ИОРы с 2025-01-01 по 2025-01-31').intervals,[('2025-01-01','2025-02-01')])
        self.assertEqual(parse_period('1 и 3 кварталы 2025').intervals,parse_period('Q1 и Q3 2025').intervals)
        self.assertEqual(parse_period('1-3 кварталы 2025').intervals,parse_period('Q1-Q3 2025').intervals)
        self.assertEqual(self.select('Черновики ИОР за 2025').incdnt_sid.tolist(),[])
        for prompt in ['ИОРы за Q5 2025','ИОРы за 5 квартал 2025']:
            with self.assertRaises(ClarificationRequired): self.select(prompt)

class FlowTests(unittest.IsolatedAsyncioTestCase):
    async def test_second_pass_followup_uses_latest_extract(self):
        with patch.object(reports,'get_session_extract',return_value={'df':frame()}), \
             patch.object(reports,'search_small_index',return_value=['ИОР EVE-1: Первый']) as search, \
             patch.object(reports,'answer_follow_up_with_qwen',return_value='Подробности') as answer, \
             patch.object(reports,'get_data_store') as store:
            result=await reports.run_ior_report(None,'latest-extract','Расскажи подробнее про EVE-1 из этой выборки')
            self.assertEqual(result,'Подробности')
            search.assert_called_once(); answer.assert_called_once(); store.assert_not_called()

    async def test_second_pass_legacy_safe_error_and_normal_clarification(self):
        import importlib.util
        spec=importlib.util.spec_from_file_location('ior_legacy_compat_test',Path(reports.__file__).parent/'tool.py')
        legacy=importlib.util.module_from_spec(spec); spec.loader.exec_module(legacy)
        tool=legacy.IORAnalyzerTool()
        with patch.object(legacy,'run_ior_report',side_effect=RuntimeError('private DSN/password')), \
             self.assertLogs(legacy.logger,level='ERROR') as logged:
            result=await tool.execute('ИОРы')
        self.assertNotIn('private',result); self.assertNotIn('Traceback',result)
        self.assertIn('журнал',result)
        self.assertIn('Traceback',' '.join(logged.output))
        with patch.object(legacy,'run_ior_report',new=AsyncMock(return_value='Уточните блок')):
            self.assertEqual(await tool.execute('ИОРы'),'Уточните блок')
        for phrase in ['SBR','приоритет','одного EVE','Уточнение:','вкладк']:
            self.assertIn(phrase,tool.description)
    async def test_native_clarification_is_normal_string(self):
        from workspace.tools.ior_analyzer import IORAnalyzerTool,IORAnalyzerToolConfig
        tool=IORAnalyzerTool(config=IORAnalyzerToolConfig())
        with patch.object(IORAnalyzerTool,'_load_runner',return_value=reports.run_ior_report),patch.object(ground,'load_catalog',return_value=CAT),patch.object(reports,'get_data_store') as store:
            result=await tool.execute(prompt='ИОРы по блоку Риски',session_id='native-clarify')
            self.assertIsInstance(result,str)
            self.assertIn('возникли',result)
            store.assert_not_called()

    async def test_internal_sort_without_limit_preserves_population(self):
        from types import SimpleNamespace
        from utils.tools import dataframe_ops
        raw=pd.DataFrame({'value':list(range(20))})
        saved=[]
        def register(df,**kwargs):
            saved.append(df)
            return SimpleNamespace(df_id='sorted',rows=len(df))
        ctx=SimpleNamespace(dataframes={'input':raw},get_df=lambda key:raw,register_dataframe=register)
        result=await dataframe_ops.top_n(ctx,'input','value',n=None)
        self.assertTrue(result.ok,result.error)
        self.assertEqual(len(saved[0]),20)
    async def test_clarification_has_no_side_effects(self):
        with patch.object(ground,'load_catalog',return_value=CAT), patch.object(reports,'get_data_store') as store, patch.object(reports,'generate_hypothesis_narrative') as narrative:
            result=await reports.run_ior_report(None,'clarify','ИОРы по блоку Риски')
            self.assertIn('возникли',result)
            store.assert_not_called(); narrative.assert_not_called()
    async def test_new_requests_ignore_old_index(self):
        with patch.object(ground,'load_catalog',return_value=CAT),patch.object(reports,'get_session_extract',return_value={'df':frame()}),patch.object(reports,'search_small_index') as search,patch.object(reports,'get_data_store',side_effect=RuntimeError('fresh query requested')):
            for prompt in ['Найди ИОРы по СРБ за 2025 год','Найди EVE-123']:
                with self.assertRaisesRegex(RuntimeError,'fresh query'): await reports.run_ior_report(None,'new',prompt)
            search.assert_not_called()
    async def test_explicit_followup(self):
        with patch.object(reports,'get_session_extract',return_value={'df':frame()}),patch.object(reports,'search_small_index',return_value=['EVE-1: Первый']) as search,patch.object(reports,'answer_follow_up_with_qwen',return_value='Ответ'),patch.object(reports,'get_data_store') as store:
            self.assertEqual(await reports.run_ior_report(None,'s','Найди EVE-1 в этой выборке'),'Ответ')
            search.assert_called_once(); store.assert_not_called()
    async def test_hypothesis_population(self):
        import ior_hypothesis as h
        for count,expected_calls in [(49,0),(50,1),(1,0)]:
            df=pd.concat([frame().iloc[:1]]*count,ignore_index=True)
            df['incdnt_sid']=['EVE-'+str(i) for i in range(count)]
            if count==1: df=pd.concat([df]*100,ignore_index=True)
            with patch.object(h,'analyze_incident_descriptions',new=AsyncMock(return_value='')) as evidence,patch.object(h,'ask_local_qwen',side_effect=RuntimeError('LLM unavailable')):
                await h.generate_hypothesis_narrative('ИОРы',df,session_id='threshold',preset_name='ior_hypothesis')
                self.assertEqual(evidence.call_count,expected_calls)
        with patch.object(h,'analyze_incident_descriptions',new=AsyncMock()) as evidence,patch.object(h,'ask_local_qwen') as llm:
            await h.generate_hypothesis_narrative('Досье EVE-1',frame().iloc[:1],session_id='dossier',preset_name='report_period_specific_ior')
            evidence.assert_not_called(); llm.assert_not_called()

class IndexTests(unittest.TestCase):
    def test_versions_and_failure(self):
        session='version-test'
        first=set_session_extract(session,frame())
        with patch.object(bge.osiris_client,'build_index',return_value=True):
            self.assertTrue(bge.build_and_cache_small_index(session,frame()))
        second=set_session_extract(session,frame().iloc[:0])
        self.assertNotEqual(first['version'],second['version'])
        self.assertEqual(bge.search_small_index(session,'query'),[])
        set_session_extract(session,frame())
        with patch.object(bge.osiris_client,'build_index',side_effect=RuntimeError('OSIRIS unavailable')):
            self.assertFalse(bge.build_and_cache_small_index(session,frame()))
        self.assertNotIn(session,bge._SMALL_FAISS_SESSION_CACHE)
    def test_search_identity_and_rerank(self):
        session='search-version'
        extract=set_session_extract(session,frame())
        with patch.object(bge.osiris_client,'build_index',return_value=True): bge.build_and_cache_small_index(session,frame())
        with patch.object(bge.osiris_client,'search',return_value=[{'id':'EVE-1','text':'Первый','score':.9}]) as search,patch.object(bge.osiris_client,'rerank',return_value=[{'id':'EVE-1','score':.8}]) as rerank:
            self.assertEqual(bge.search_small_index(session,'найди'),['ИОР EVE-1: Первый'])
            self.assertEqual(search.call_args.args[1],extract['version']); rerank.assert_called_once()
        bge._SMALL_FAISS_SESSION_CACHE[session]['version']='old'
        self.assertEqual(bge.search_small_index(session,'query'),[])

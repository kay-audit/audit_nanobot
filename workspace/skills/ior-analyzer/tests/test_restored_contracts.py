"""Business acceptance on real, isolated DuckDB tables; no live services."""
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import AsyncMock, Mock

import duckdb
import openpyxl
import pandas as pd
import pytest
import yaml

import ior_reports as reports
from utils.data_store import DUCKDB_TABLES
from utils.ior_artifacts import artifact_scope
from utils.resolve import semantic_grounding as grounding
from utils.resolve.request_plan import resolve_preset, build_plan
from utils.resolve.request_outcome import ClarificationRequired
from preset_analysis.report_period_specific_ior_analysis import prepare_dossier_views

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = Path(__file__).with_name('report_field_contracts.json')
BIG_ID = 9007199254740993
PROMPTS = [
    ('ИОРы за март 2025 года по СРБ', 'ior_hypothesis'),
    ('Покажи все ИОР за 2025 год', 'ior_hypothesis'),
    ('Выгрузи ИОР по блоку Риски', 'ior_hypothesis'),
    ('Сколько всего ИОР и какова общая сумма потерь?', 'ior_hypothesis'),
    ('Покажи ИОР по ЦПР DRP-10047', 'ior_hypothesis'),
    ('Финансовые последствия ИОР за март', 'financial_consequences_ior'),
    ('Покажи виды финансовых потерь', 'financial_consequences_ior'),
    ('Покажи возмещения по ИОР', 'vozmeshenie_ior'),
    ('Найди удалённые ИОР с причинами удаления', 'deleted_ior'),
    ('ИОР с репутационными последствиями', 'ior_nonfinancial_consequences'),
    ('Нефинансовые последствия с высоким влиянием', 'ior_nonfinancial_consequences'),
    ('Полное досье EVE-5092355', 'report_period_specific_ior'),
    ('Финансовые последствия EVE-5092355', 'financial_consequences_ior'),
    ('Все ИОР Сбербанка', 'ior_hypothesis'),
    ('Общая сумма ущерба', 'ior_hypothesis'),
]


@pytest.mark.parametrize('hint', [None, 'ior_period_pao_sberbank', 'financial_consequences_ior'])
@pytest.mark.parametrize('prompt,expected', PROMPTS)
def test_user_intent_wins_over_llm_hint(hint, prompt, expected):
    assert resolve_preset(hint, prompt) == expected


def test_trusted_technical_selection_and_ambiguity():
    assert resolve_preset('ior_period_pao_sberbank', 'ИОРы за март по СРБ', trusted_preset=True) == 'ior_period_pao_sberbank'
    assert resolve_preset('vozmeshenie_ior', 'Общая статистика', trusted_preset=True) == 'vozmeshenie_ior'
    assert resolve_preset('credit_no_way_collect_debt', '', trusted_preset=True) == 'credit_no_way_collect_debt'
    with pytest.raises(ClarificationRequired):
        resolve_preset('financial_consequences_ior', 'Возмещения и финансовые последствия')
    with pytest.raises(ClarificationRequired):
        resolve_preset(None, 'Финансовые и нефинансовые последствия')


class Store:
    tables = DUCKDB_TABLES
    backend_name = 'isolated_duckdb_fixture'

    def __init__(self):
        self.connection = duckdb.connect(':memory:')
        self.queries = []

    def query_sql(self, sql):
        self.queries.append(sql)
        return self.connection.execute(sql).df()


@pytest.fixture
def store(monkeypatch):
    store = Store()
    schema = yaml.safe_load((ROOT / 'utils/schema/kb_schema.yaml').read_text(encoding='utf-8'))['tables']
    for logical in ('ior', 'financial_impact', 'recovery', 'status', 'nonfinancial_impact'):
        table = DUCKDB_TABLES[logical]
        columns = schema[table]['columns']
        def dtype(column):
            t = column['type'].lower()
            if t.startswith('decimal'): return 'DECIMAL(18,4)'
            if t in ('timestamp', 'date', 'bigint', 'int', 'integer', 'boolean'): return t
            return 'VARCHAR'
        store.connection.execute('CREATE TABLE ' + table + ' (' + ', '.join('"' + c['name'] + '" ' + dtype(c) for c in columns) + ')')

    def insert(logical, **record):
        store.connection.execute('INSERT INTO ' + DUCKDB_TABLES[logical] + ' (' + ','.join(record) + ') VALUES (' + ','.join('?' for _ in record) + ')', list(record.values()))

    for i, (sid, day, bank, status) in enumerate([
        ('EVE-5092355', '2025-03-01', 'Среднерусский банк', 'Утверждён'),
        ('EVE-2', '2025-04-01', 'Среднерусский банк', 'Утверждён'),
        ('EVE-3', '2025-03-31', 'Московский банк', 'Утверждён'),
        ('EVE-4', '2025-03-20', 'Среднерусский банк', 'Удалён'),
        ('EVE-5', '2025-03-21', 'Среднерусский банк', 'Удален'),
        ('EVE-6', '2025-03-15', 'Среднерусский банк', 'Утверждён'),
    ]):
        insert('ior', incdnt_id=BIG_ID+i, incdnt_sid=sid, incdnt_entry_dt=day,
               org_struct_lvl_3_name=bank, org_struct_lvl_4_name='Другое подразделение',
               funct_block_lvl_3_name='Риски', risk_profile_id='DRP-10047',
               incdnt_status_name=status, org_struct_id='SBR_A',
               incdnt_full_descr_txt='=Описание без формулы', incdnt_sum=Decimal('9999.25'))
    for i, amount in enumerate(['100.25', '200.50', '300.75']):
        insert('financial_impact', incdnt_id=BIG_ID, fin_impact_id=BIG_ID+1000+i,
               fin_impact_sid='FIN-'+str(i), fin_impact_type_name='Финансовое последствие',
               fin_impact_kind_name='Потеря', fin_impact_rub_amt=Decimal(amount),
               fin_impact_detection_dt='2025-03-02', fin_impact_account_num='0001234567890123456789')
    for i, amount in enumerate(['10.25', '20.50']):
        insert('recovery', incdnt_id=BIG_ID, recovery_sid='REC-'+str(i),
               recovery_type_name='Возмещение', recovery_rub_amt=Decimal(amount), recovery_reg_dt='2025-03-03')
        insert('nonfinancial_impact', incdnt_id=BIG_ID, nonfin_impact_sid='NFI-'+str(i),
               nonfin_impact_kind_name='Репутационный риск', nonfin_impact_influence_class_name='Высокий')
    for i, code in enumerate(['delete', 'approveDeletion', 'rejectDeletion']):
        insert('status', incdnt_id=BIG_ID+3, incdnt_status_name='Удалён',
               stts_chng_action_code=code, stts_chng_action_name='Согласование',
               stts_chng_action_dttm='2025-04-0'+str(i+1), stts_chng_user_num='000017',
               stts_chng_comment_txt='Причина '+str(i))
    monkeypatch.setattr(grounding, 'load_catalog', lambda: {
        grounding.ORIGIN: ['Среднерусский банк', 'Московский банк'],
        grounding.RESPONSIBILITY: ['Риски'],
    })
    yield store
    store.connection.close()


CASES = [
    ('ior_hypothesis', 'Все ИОР за 2025 год', 6, 6),
    ('financial_consequences_ior', 'Финансовые последствия EVE-5092355', 3, 1),
    ('vozmeshenie_ior', 'Возмещения EVE-5092355', 2, 1),
    ('deleted_ior', 'Удалённые ИОР за март 2025', 3, 2),
    ('ior_nonfinancial_consequences', 'Нефинансовые последствия EVE-5092355', 2, 1),
    ('report_period_specific_ior', 'Полное досье EVE-5092355', 6, 1),
]


@pytest.mark.parametrize('preset,prompt,rows,incidents', CASES)
@pytest.mark.asyncio
async def test_complete_pipeline_xlsx(store, monkeypatch, tmp_path, preset, prompt, rows, incidents):
    fields = json.loads(MANIFEST.read_text(encoding='utf-8'))[preset]
    monkeypatch.setattr(reports, 'get_data_store', lambda: store)
    monkeypatch.setattr(reports, 'get_session_extract', lambda _: None)
    monkeypatch.setattr(reports, 'build_and_cache_small_index', Mock())
    narrative = AsyncMock(return_value='Детерминированная проверка без LLM')
    monkeypatch.setattr(reports, 'generate_hypothesis_narrative', narrative)
    paths = []
    with artifact_scope(tmp_path, paths):
        await reports.run_ior_report('ior_period_pao_sberbank', 'fixture-'+preset, prompt, trusted_preset=False)
    assert len(paths) == 1
    raw = narrative.call_args.args[1]
    assert list(raw.columns) == [f['technical_name'] for f in fields]
    assert len(raw) == rows and raw.incdnt_id.nunique() == incidents
    assert narrative.call_args.kwargs['preset_name'] == preset
    assert all('LIMIT' not in q.upper() for q in store.queries)
    workbook = openpyxl.load_workbook(paths[0], data_only=False)
    try:
        sheet = workbook.active
        values = list(sheet.values)
        headers = values[0]
        assert headers == tuple(f['excel_label'] for f in fields)
        assert len(set(headers)) == len(headers)
        assert len(values)-1 == rows
        assert all(isinstance(v[0], str) for v in values[1:])
        expected_id = BIG_ID+3 if preset == 'deleted_ior' else BIG_ID
        assert str(expected_id) in {r[0] for r in values[1:]}
        assert all(any('\u0400' <= c <= '\u04ff' for c in h) for h in headers)
        by_name = {f['technical_name']: i for i, f in enumerate(fields)}
        assert isinstance(values[1][by_name['incdnt_entry_dt']], datetime)
        assert values[1][by_name['incdnt_mistake_cnt']] is None
        description = sheet.cell(2, by_name['incdnt_full_descr_txt']+1)
        assert description.value == '=Описание без формулы' and description.data_type == 's'
        if 'fin_impact_id' in by_name:
            assert str(BIG_ID+1000) in {r[by_name['fin_impact_id']] for r in values[1:]}
            assert values[1][by_name['fin_impact_account_num']] == '0001234567890123456789'
            assert sum(raw.fin_impact_rub_amt.drop_duplicates()) == pytest.approx(601.5)
            assert {r[by_name['fin_impact_rub_amt']] for r in values[1:]} == {100.25, 200.5, 300.75}
        if 'recovery_rub_amt' in by_name:
            assert {r[by_name['recovery_rub_amt']] for r in values[1:]} == {10.25, 20.5}
        if 'incdnt_drct_dmg_sum' in by_name:
            assert values[1][by_name['incdnt_drct_dmg_sum']] is None
        if preset == 'report_period_specific_ior':
            metrics = prepare_dossier_views(raw)
            assert (metrics['raw_rows'], metrics['fin_count'], metrics['recovery_count']) == (6, 3, 2)
            assert metrics['fin_total'] == pytest.approx(601.5)
            assert metrics['recovery_total'] == pytest.approx(30.75)
        if preset == 'deleted_ior':
            assert set(raw.stts_chng_action_code.dropna()) == {'delete', 'approveDeletion'}
            assert raw.loc[raw.incdnt_sid == 'EVE-5', 'stts_chng_action_code'].isna().all()
    finally:
        workbook.close()


def test_srb_march_and_dossier_without_children(store):
    plan = build_plan(None, 'ИОРы за март 2025 года по СРБ DRP-10047')
    assert plan.period.intervals == [('2025-03-01', '2025-04-01')]
    sql = reports.build_dynamic_sql_from_prompt('ИОРы за март 2025 года по СРБ DRP-10047', tables=store.tables)
    result = store.query_sql(sql)
    assert set(result.incdnt_sid) == {'EVE-5092355', 'EVE-4', 'EVE-5', 'EVE-6'}
    assert set(result.org_struct_lvl_3_name) == {'Среднерусский банк'}
    assert 'DRP-10047' in sql
    raw = store.query_sql(reports.build_dynamic_sql_from_prompt('Полное досье EVE-6', tables=store.tables))
    assert len(raw) == 1 and raw.fin_impact_sid.isna().all() and raw.recovery_sid.isna().all()
    export = reports.prepare_report_export(raw, 'report_period_specific_ior')
    assert len(export) == 1


def test_schema_variation_preserves_actual_fields(store):
    store.connection.execute('ALTER TABLE '+store.tables['ior']+' DROP COLUMN clntpth_lvl_4_name')
    store.connection.execute('ALTER TABLE '+store.tables['ior']+' ADD COLUMN created_at TIMESTAMP')
    from preset_analysis.report_columns import ordered_main_columns
    columns = ordered_main_columns(store.query_sql('SELECT * FROM '+store.tables['ior']+' WHERE 1=0').columns)
    raw = store.query_sql(reports.build_preset_sql_queries(store.tables, columns)['ior_hypothesis'])
    assert 'clntpth_lvl_4_name' not in raw and 'created_at' in raw
    assert list(raw.columns) == list(columns)
    with pytest.raises(ValueError, match='required fields'):
        ordered_main_columns(['incdnt_id'])


def test_main_duplicates_do_not_multiply_aggregates(store):
    raw = store.query_sql(reports.build_preset_sql_queries(store.tables)['ior_hypothesis'])
    repeated = pd.concat([raw, raw], ignore_index=True)
    out = reports.prepare_report_export(repeated, 'ior_hypothesis')
    assert len(out) == 6
    assert out[reports.CYRILLIC_RENAME['incdnt_sum']].sum() == pytest.approx(6*9999.25)


def test_atomic_openpyxl_fallback_keeps_literal_text_and_null(monkeypatch, tmp_path):
    from utils.dataframe_ops import _write_xlsx_safe
    original = pd.DataFrame.to_excel

    def without_xlsxwriter(self, *args, **kwargs):
        if kwargs.get('engine') == 'xlsxwriter':
            raise ImportError('fixture: optional xlsxwriter unavailable')
        return original(self, *args, **kwargs)

    monkeypatch.setattr(pd.DataFrame, 'to_excel', without_xlsxwriter)
    path = tmp_path / 'fallback.xlsx'
    _write_xlsx_safe(pd.DataFrame({'Описание': ['=literal'], 'Сумма': [None]}), path, 'ИОР', format_columns=False)
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        assert workbook['ИОР']['A2'].data_type == 's'
        assert workbook['ИОР']['A2'].value == '=literal'
        assert workbook['ИОР']['B2'].value is None
        assert not path.with_suffix('.xlsx.tmp').exists()
    finally:
        workbook.close()


@pytest.mark.asyncio
async def test_empty_filtered_population_never_queries_wider_source(store, monkeypatch):
    monkeypatch.setattr(reports, 'get_data_store', lambda: store)
    monkeypatch.setattr(reports, 'get_session_extract', lambda _: None)
    narrative = AsyncMock(return_value='Нет совпадений')
    monkeypatch.setattr(reports, 'generate_hypothesis_narrative', narrative)
    index = Mock()
    monkeypatch.setattr(reports, 'build_and_cache_small_index', index)
    result = await reports.run_ior_report(None, 'empty-real-fixture', 'Возмещения EVE-999999 за март 2025')
    assert narrative.call_args.args[1].empty
    assert 'Источник данных доступен' in result
    assert len(store.queries) == 2  # schema-only probe + one filtered business query
    index.assert_not_called()
